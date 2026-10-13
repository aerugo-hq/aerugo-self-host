"""Document formats (slices L and L2): one extraction path for every connector.

Markdown, plain text, HTML, Word, PowerPoint, PDF, spreadsheets (Excel and CSV), OpenDocument (text,
spreadsheets, and presentations), and email all become (title, sections), where a section is (heading path,
text): the same shape the Markdown chunker produces, so everything downstream (chunks, embeddings, bronze
concepts, search) is unchanged. Google Docs arrive as Word exports. Tables become sections of 50 rows that each
repeat the header row, so a match deep in a sheet still says what its columns mean; past 5,000 rows a sheet is
indexed in part, and says so. Spreadsheets and OpenDocument files are read with the standard library alone.
Scanned PDFs and pictures of text are read with OCR (Tesseract, with Poppler rendering PDF pages), and RTF with
a small converter (slice L2b).

A file that can't be read raises ExtractionError with a reason. Nothing is ever indexed as empty:
malformed files, encrypted PDFs, scans OCR finds no text in, images with only a few words (a logo, a diagram),
and archives that would inflate past a safe size are all refused, visibly.
"""
import csv
import io
import logging
import os
import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from html.parser import HTMLParser

from .markdown import parse as parse_markdown

# pypdf logs warnings about damaged files; those files are refused with an ExtractionError instead.
logging.getLogger("pypdf").setLevel(logging.ERROR)

TEXT_FORMATS = (".md", ".markdown", ".html", ".htm", ".txt", ".csv", ".eml", ".rtf")
IMAGE_FORMATS = (".png", ".jpg", ".jpeg", ".tif", ".tiff")
BINARY_FORMATS = (".docx", ".pptx", ".pdf", ".xlsx", ".odt", ".ods", ".odp") + IMAGE_FORMATS
SUPPORTED = TEXT_FORMATS + BINARY_FORMATS
MAX_UNZIPPED = 100 * 1024 * 1024  # Office and OpenDocument files are zips; refuse ones that inflate past this
MAX_PDF_PAGES = 2000
MAX_TABLE_ROWS = 5000  # rows indexed from one sheet or CSV; a larger one is indexed in part, and says so
ROWS_PER_SECTION = 50
OCR_MAX_PAGES = 50       # pages of a scanned PDF read with OCR; past them, the rest is noted
OCR_TIMEOUT = 120        # seconds for one page or image
OCR_MIN_WORDS = 5        # an image with fewer words than this (a logo, a diagram) isn't indexed
MAX_IMAGE_PIXELS = 50_000_000


class ExtractionError(ValueError):
    """The file can't be turned into text. The message says why, and is safe to show and log."""


def supported(path: str) -> bool:
    return path.lower().endswith(SUPPORTED)


def extract(path: str, data: bytes, fallback_title: str) -> tuple[str, list[tuple[str | None, str]]]:
    name = path.lower()
    if name.endswith((".md", ".markdown")):
        markdown = data.decode("utf-8", errors="replace")
    elif name.endswith((".html", ".htm")):
        markdown = _html(data.decode("utf-8", errors="replace"))
    elif name.endswith(".docx"):
        markdown = _docx(data)
    elif name.endswith(".pptx"):
        markdown = _pptx(data)
    elif name.endswith(".pdf"):
        markdown = _pdf(data)
    elif name.endswith(".txt"):
        markdown = data.decode("utf-8", errors="replace")
    elif name.endswith(".csv"):
        markdown = _csv(data)
    elif name.endswith(".xlsx"):
        markdown = _xlsx(data)
    elif name.endswith((".odt", ".ods", ".odp")):
        markdown = _odf(data)
    elif name.endswith(".eml"):
        markdown = _eml(data)
    elif name.endswith(".rtf"):
        markdown = _rtf(data)
    elif name.endswith(IMAGE_FORMATS):
        markdown = _image(data, name.rsplit(".", 1)[-1])
    else:
        raise ExtractionError(f"unsupported format: {path.rsplit('.', 1)[-1]}")
    title, sections = parse_markdown(markdown, fallback_title=fallback_title)
    if not any(text.strip() for _, text in sections):
        raise ExtractionError("no text found in the file")
    return title, sections


def _guard_zip(data: bytes) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise ExtractionError("not a valid Office file (the archive is damaged or isn't one)")
    if sum(info.file_size for info in archive.infolist()) > MAX_UNZIPPED:
        raise ExtractionError("the file would inflate past the size limit, so it wasn't opened")
    return archive


def _cell(text: str) -> str:
    return " ".join(text.split())


def _docx(data: bytes) -> str:
    _guard_zip(data)
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    try:
        document = Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001 - python-docx raises many types for damaged files
        raise ExtractionError(f"not a readable Word document ({type(e).__name__})")
    lines = []
    for child in document.element.body.iterchildren():  # paragraphs and tables, in document order
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            if style == "Title":
                lines.append(f"# {text}")
            elif style.startswith("Heading ") and style[8:].isdigit():
                lines.append(f"{'#' * min(int(style[8:]) + 1, 6)} {text}")  # Heading 1 sits under the title
            else:
                lines.append(text)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                lines.append(" | ".join(_cell(c.text) for c in row.cells))
        lines.append("")
    return "\n".join(lines)


def _pptx(data: bytes) -> str:
    _guard_zip(data)
    from pptx import Presentation
    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001 - python-pptx raises many types for damaged files
        raise ExtractionError(f"not a readable PowerPoint file ({type(e).__name__})")
    lines = []
    for number, slide in enumerate(deck.slides, 1):
        title_shape = slide.shapes.title
        title = (title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else "")
        lines.append(f"## {title or f'Slide {number}'}")
        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if shape.has_text_frame:
                lines.extend(p.strip() for p in shape.text_frame.text.splitlines() if p.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                lines.extend(" | ".join(_cell(c.text) for c in row.cells) for row in shape.table.rows)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if notes:
            lines.append(f"Speaker notes: {notes}")
        lines.append("")
    return "\n".join(lines)


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("the PDF is encrypted")
        pages = reader.pages
        if len(pages) > MAX_PDF_PAGES:
            raise ExtractionError(f"the PDF has more than {MAX_PDF_PAGES} pages")
        texts = [(page.extract_text() or "").strip() for page in pages]
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError) as e:
        raise ExtractionError(f"not a readable PDF ({type(e).__name__})")
    if not any(texts):
        return _ocr_pdf(data, len(pages))
    return "\n\n".join(f"## Page {n}\n\n{text}" for n, text in enumerate(texts, 1) if text)


class _HTMLToMarkdown(HTMLParser):
    SKIP = {"script", "style", "noscript", "nav", "footer", "header", "template", "svg"}
    BLOCKS = {"p", "div", "li", "tr", "br", "section", "article", "blockquote", "pre", "dd", "dt"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.lines: list[str] = []
        self.current: list[str] = []
        self.skipping = 0
        self.heading = 0
        self.title = ""
        self.in_title = False
        self.row: list[str] | None = None

    def _flush(self):
        text = " ".join("".join(self.current).split())
        if text:
            self.lines.append(f"{'#' * self.heading} {text}" if self.heading else text)
        self.current = []

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skipping += 1
        elif tag == "title":
            self.in_title = True
        elif tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._flush()
            self.heading = int(tag[1])
        elif tag == "tr":
            self._flush()
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            self._flush()
        elif tag in self.BLOCKS:
            self._flush()

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skipping = max(0, self.skipping - 1)
        elif tag == "title":
            self.in_title = False
        elif tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._flush()
            self.heading = 0
        elif tag in {"td", "th"} and self.row is not None:
            self.row.append(" ".join("".join(self.current).split()))
            self.current = []
        elif tag == "tr" and self.row is not None:
            if any(self.row):
                self.lines.append(" | ".join(self.row))
            self.row = None
        elif tag in self.BLOCKS:
            self._flush()

    def handle_data(self, data):
        if self.in_title:
            self.title += data
        elif not self.skipping:
            self.current.append(data)


def _html(html: str) -> str:
    parser = _HTMLToMarkdown()
    parser.feed(html)
    parser.close()
    parser._flush()
    body = "\n\n".join(parser.lines)
    title = " ".join(parser.title.split())
    if title and not body.lstrip().startswith("# "):
        body = f"# {title}\n\n{body}"
    return body


# ------------------------------------------------------------------ tables, spreadsheets, OpenDocument, email (slice L2)
def _clean_rows(rows: list[list]) -> list[list[str]]:
    return [[_cell(str(c)).replace("|", "\\|") for c in row] for row in rows if any(str(c).strip() for c in row)]


def _markdown_rows(header: list[str], body: list[list[str]]) -> str:
    width = max(len(r) for r in [header, *body])
    pad = lambda r: r + [""] * (width - len(r))  # noqa: E731
    return "\n".join(["| " + " | ".join(pad(header)) + " |", "|" + " --- |" * width] +
                     ["| " + " | ".join(pad(r)) + " |" for r in body])


def _inline_table(rows: list[list]) -> str:
    """A table inside a document: Markdown rows under whatever heading it sits in, as Word tables are."""
    rows = _clean_rows(rows)
    return _markdown_rows(rows[0], rows[1:MAX_TABLE_ROWS + 1]) if rows else ""


def _table(name: str, rows: list[list[str]], total: int | None = None) -> str:
    """A sheet as Markdown sections of ROWS_PER_SECTION rows, each repeating the header row, so a chunk stands on
    its own. The first non-empty row is the header. Past MAX_TABLE_ROWS, the rest is noted rather than dropped."""
    rows = _clean_rows(rows)
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    total = max(total or 0, len(body))
    body = body[:MAX_TABLE_ROWS]
    sections = []
    for start in range(0, max(len(body), 1), ROWS_PER_SECTION):
        block = body[start:start + ROWS_PER_SECTION]
        label = f"{name}, rows {start + 1}\u2013{start + len(block)}" if block else name
        sections.append(f"## {label}\n\n" + _markdown_rows(header, block))
    if total > len(body):
        sections.append(f"## {name}, the rest\n\nThe first {len(body):,} of {total:,} rows are indexed.")
    return "\n\n".join(sections)


def _csv(data: bytes) -> str:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows, total = [], 0
    for row in csv.reader(io.StringIO(text), dialect):
        total += 1
        if len(rows) <= MAX_TABLE_ROWS:
            rows.append(row)
    return _table("Rows", rows, total - 1)


_SHEET = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _column(ref: str) -> int | None:
    """A cell reference's column, from zero: "A1" is 0, "AB12" is 27."""
    letters = "".join(ch for ch in ref if ch.isalpha()).upper()
    if not letters:
        return None
    index = 0
    for ch in letters:
        index = index * 26 + (ord(ch) - 64)
    return index - 1


def _xlsx(data: bytes) -> str:
    archive = _guard_zip(data)
    try:
        names = set(archive.namelist())
        shared = []
        if "xl/sharedStrings.xml" in names:
            for item in ET.fromstring(archive.read("xl/sharedStrings.xml")).iter(f"{_SHEET}si"):
                shared.append("".join(t.text or "" for t in item.iter(f"{_SHEET}t")))
        targets = {r.get("Id"): r.get("Target", "") for r in ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))}
        out = []
        for sheet in ET.fromstring(archive.read("xl/workbook.xml")).iter(f"{_SHEET}sheet"):
            target = targets.get(sheet.get(f"{_REL}id"), "")
            path = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
            rows = []
            for row in ET.fromstring(archive.read(path)).iter(f"{_SHEET}row"):
                cells: dict[int, str] = {}
                for c in row.iter(f"{_SHEET}c"):
                    kind, v = c.get("t"), c.find(f"{_SHEET}v")
                    if kind == "inlineStr":
                        value = "".join(t.text or "" for t in c.iter(f"{_SHEET}t"))
                    else:
                        value = "" if v is None or v.text is None else v.text
                        if kind == "s" and value:
                            value = shared[int(value)]
                        elif kind == "b":
                            value = "TRUE" if value == "1" else "FALSE"
                    col = _column(c.get("r", ""))
                    cells[len(cells) if col is None else col] = value
                if cells:
                    rows.append([cells.get(i, "") for i in range(max(cells) + 1)])
            table = _table(sheet.get("name") or "Sheet", rows)
            if table:
                out.append(table)
        return "\n\n".join(out)
    except (KeyError, ET.ParseError, ValueError, IndexError) as e:
        raise ExtractionError(f"not a readable spreadsheet ({type(e).__name__})")


_ODF = {"office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0", "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
        "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0", "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"}


def _odf(data: bytes) -> str:
    """OpenDocument text, spreadsheets, and presentations: headings at their level (level 1 under the title, as
    for Word), paragraphs, tables (each spreadsheet sheet in sections of its own, a table in a document inline under
    its heading), and each slide under its name."""
    archive = _guard_zip(data)
    try:
        body = ET.fromstring(archive.read("content.xml")).find(f"{{{_ODF['office']}}}body")
    except (KeyError, ET.ParseError) as e:
        raise ExtractionError(f"not a readable OpenDocument file ({type(e).__name__})")
    if body is None:
        raise ExtractionError("not a readable OpenDocument file (no body)")
    tag = lambda ns, name: f"{{{_ODF[ns]}}}{name}"  # noqa: E731
    spreadsheet = body.find(tag("office", "spreadsheet")) is not None
    lines: list[str] = []

    def text_of(el) -> str:
        return _cell("".join(el.itertext()))

    def walk(el) -> None:
        if el.tag == tag("text", "h"):
            level = int(el.get(tag("text", "outline-level")) or 1)
            lines.append(f"{'#' * min(level + 1, 6)} {text_of(el)}")
        elif el.tag == tag("text", "p"):
            if text_of(el):
                lines.append(text_of(el))
        elif el.tag == tag("table", "table"):
            rows = [[text_of(cell) for cell in row.iter(tag("table", "table-cell"))] for row in el.iter(tag("table", "table-row"))]
            lines.append(_table(el.get(tag("table", "name")) or "Sheet", rows) if spreadsheet else _inline_table(rows))
        elif el.tag == tag("draw", "page"):
            lines.append(f"## {el.get(tag('draw', 'name')) or 'Slide'}")
            for child in el:
                walk(child)
        else:
            for child in el:
                walk(child)

    walk(body)
    return "\n\n".join(line for line in lines if line)


def _eml(data: bytes) -> str:
    """An email: its subject as the title, who sent it to whom and when, its text (or its HTML, as text), and the
    names of its attachments, which aren't opened."""
    from email import policy
    from email.parser import BytesParser
    try:
        message = BytesParser(policy=policy.default).parsebytes(data)
        part = message.get_body(preferencelist=("plain", "html"))
        text = "" if part is None else part.get_content()
        if part is not None and part.get_content_type() == "text/html":
            text = _html(text)
        attachments = [a.get_filename() for a in message.iter_attachments() if a.get_filename()]
    except Exception as e:  # noqa: BLE001 - the email package raises many types for damaged messages
        raise ExtractionError(f"not a readable email ({type(e).__name__})")
    header = "\n".join(f"{name}: {_cell(str(message[name]))}" for name in ("From", "To", "Cc", "Date") if message[name])
    parts = [f"# {_cell(str(message['Subject'] or '')) or 'Email'}", header, text.strip()]
    if attachments:
        parts.append("Attachments (not opened): " + ", ".join(attachments))
    return "\n\n".join(p for p in parts if p)


# ------------------------------------------------------------------ OCR and RTF (slice L2b)
def _tesseract(path: str) -> str:
    try:
        done = subprocess.run(["tesseract", path, "-", "--psm", "3"], capture_output=True, timeout=OCR_TIMEOUT, check=False)
    except subprocess.TimeoutExpired:
        raise ExtractionError("OCR took too long on a page, so the file wasn't indexed")
    return done.stdout.decode("utf-8", errors="replace").strip() if done.returncode == 0 else ""


def _ocr_pdf(data: bytes, page_count: int) -> str:
    """A scanned PDF: render its pages (up to OCR_MAX_PAGES) and read each with Tesseract."""
    if not (shutil.which("tesseract") and shutil.which("pdftoppm")):
        raise ExtractionError("the PDF has no text layer (a scan?), and OCR isn't available here")
    with tempfile.TemporaryDirectory() as tmp:
        source = os.path.join(tmp, "scan.pdf")
        with open(source, "wb") as f:
            f.write(data)
        try:
            subprocess.run(["pdftoppm", "-r", "200", "-gray", "-png", "-l", str(OCR_MAX_PAGES), source,
                            os.path.join(tmp, "page")], capture_output=True, timeout=OCR_TIMEOUT * 2, check=True)
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
            raise ExtractionError("the scanned PDF couldn't be rendered for OCR")
        pages = sorted((p for p in os.listdir(tmp) if p.startswith("page-") and p.endswith(".png")), key=lambda p: int(p[5:-4]))
        texts = [_tesseract(os.path.join(tmp, p)) for p in pages]
    if not any(texts):
        raise ExtractionError("the PDF has no text layer, and OCR found no text in it")
    out = [f"## Page {n} (read by OCR)\n\n{text}" for n, text in enumerate(texts, 1) if text]
    if page_count > OCR_MAX_PAGES:
        out.append(f"## The rest\n\nOCR read the first {OCR_MAX_PAGES} of {page_count} pages.")
    return "\n\n".join(out)


def _image(data: bytes, extension: str) -> str:
    """A picture of text (a scanned page, a whiteboard photo) read with OCR. One with only a few words, such as a
    logo or a diagram, isn't indexed; one too large to read safely isn't opened."""
    from PIL import Image, UnidentifiedImageError
    try:
        with Image.open(io.BytesIO(data)) as picture:  # reads the header only
            width, height = picture.size
    except (UnidentifiedImageError, OSError, ValueError):
        raise ExtractionError("not a readable image")
    if width * height > MAX_IMAGE_PIXELS:
        raise ExtractionError("the image is too large to read safely, so it wasn't opened")
    if not shutil.which("tesseract"):
        raise ExtractionError("reading images needs OCR, which isn't available here")
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, f"image.{extension}")
        with open(path, "wb") as f:
            f.write(data)
        text = _tesseract(path)
    if len(text.split()) < OCR_MIN_WORDS:
        raise ExtractionError("no text found in the image")
    return f"## Text in the image (read by OCR)\n\n{text}"


_RTF = re.compile(r"\\([a-z]{1,32})(-?\d{1,10})?[ ]?|\\'([0-9a-f]{2})|\\([^a-z])|([{}])|[\r\n]+|(.)", re.I | re.S)
_RTF_SKIP = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "header", "headerl", "headerr", "footer", "footerl",
             "footerr", "object", "listtable", "listoverridetable", "rsidtbl", "generator", "themedata",
             "colorschememapping", "datastore", "latentstyles", "xmlnstbl", "mmathPr"}
_RTF_TEXT = {"par": "\n", "line": "\n", "sect": "\n\n", "page": "\n\n", "tab": "\t", "emdash": "\u2014", "endash": "\u2013",
             "lquote": "\u2018", "rquote": "\u2019", "ldblquote": "\u201c", "rdblquote": "\u201d", "bullet": "\u2022"}


def _rtf(data: bytes) -> str:
    """RTF as plain text: paragraphs and characters kept, formatting, fonts, pictures, and metadata dropped."""
    source = data.decode("latin-1")
    if not source.lstrip().startswith("{\\rtf"):
        raise ExtractionError("not a readable RTF file")
    stack, ignorable, uc, skip, out = [], False, 1, 0, []
    for word, arg, hexcode, symbol, brace, char in _RTF.findall(source):
        if brace:
            skip = 0
            if brace == "{":
                stack.append((uc, ignorable))
            elif stack:
                uc, ignorable = stack.pop()
        elif symbol:
            skip = 0
            if symbol == "*":
                ignorable = True
            elif not ignorable and symbol in "{}\\":
                out.append(symbol)
            elif not ignorable and symbol == "~":
                out.append("\u00a0")
        elif word:
            skip = 0
            word = word.lower()
            if word in _RTF_SKIP:
                ignorable = True
            elif ignorable:
                continue
            elif word in _RTF_TEXT:
                out.append(_RTF_TEXT[word])
            elif word == "uc":
                uc = int(arg or 1)
            elif word == "u":
                code = int(arg or 0)
                out.append(chr(code + 65536 if code < 0 else code))
                skip = uc  # the fallback characters that follow a \u are for readers without Unicode
        elif hexcode or char:
            if skip:
                skip -= 1
            elif not ignorable:
                out.append(bytes([int(hexcode, 16)]).decode("cp1252", errors="replace") if hexcode else char)
    return "".join(out)
