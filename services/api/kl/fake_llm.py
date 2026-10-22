"""Test-only OpenAI-compatible chat model and reranker (runs only under the smoke profile).

It answers by citing source [1], drafts concepts for the compiler, grades answers for the eval, and
reranks by the words a passage shares with the query. It records every request it receives, so the
smoke suite can prove what did and didn't leave the cell: which sources reached a model or the reranker,
and that nothing reached a model at all when the caller had no visible sources. POST /mode switches its
answers to an unfaithful one (claims no source makes), so smoke can prove the grader catches it. Under /oidc it is
also a CI token issuer (slice J1): GitHub Actions-shaped workload tokens, signed with a key made at start, so smoke can
prove pipeline pushes without GitHub. It runs only under the smoke profile, so nothing trusts it otherwise.
"""
import io
import json
import os
import re
import secrets
import threading
import time
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlparse

REQUESTS: list[dict] = []
RERANKS: list[dict] = []
ALERTS: list[dict] = []  # what the app's alert webhook posted (slice P); never reset, alerts arrive late
ANSWER = "According to the sources, page the on-call first [1]."
UNFAITHFUL = "The CFO approves every refund personally [1]."  # nothing in any source says this
MODE = {"answer": ANSWER, "delay": 0.0}
# H2: any source's fake can be throttled. POST /throttle {"prefix", "count", "status", "retry_after"} answers the next
# `count` calls whose path starts with prefix with status (429 or 503), and Retry-After when given; GET /throttle says
# how many it answered so. Smoke proves a sync survives it.
THROTTLE = {"prefix": "", "left": 0, "status": 429, "retry_after": None, "answered": 0}
THROTTLE_LOCK = threading.Lock()  # delay: seconds each answer and rerank takes, when smoke asks (H5)
STOP = {"according", "sources", "source", "first", "that", "this", "with", "from", "they", "their", "there",
        "what", "when", "which", "about", "should", "would"}
ISSUER = os.environ.get("KL_FAKE_OIDC_ISSUER", "http://fake-llm:8080/oidc")
_CI_KEY: list = []


def ci_key():
    if not _CI_KEY:
        from cryptography.hazmat.primitives.asymmetric import rsa
        _CI_KEY.append(rsa.generate_private_key(public_exponent=65537, key_size=2048))
    return _CI_KEY[0]


def ci_jwks() -> dict:
    import jwt
    key = json.loads(jwt.algorithms.RSAAlgorithm.to_jwk(ci_key().public_key()))
    return {"keys": [{**key, "kid": "fake-ci", "use": "sig", "alg": "RS256"}]}


def ci_token(request: dict) -> str:
    """A workload token as GitHub Actions issues one, for the repository and ref asked for."""
    import jwt
    now, repo, ref = int(time.time()), request.get("repository", "acme/repo"), request.get("ref", "refs/heads/main")
    claims = {"iss": ISSUER, "aud": request.get("audience", "kl-ingest"), "sub": f"repo:{repo}:ref:{ref}",
              "repository": repo, "ref": ref, "iat": now, "nbf": now, "exp": now + int(request.get("ttl", 300))}
    return jwt.encode(claims, ci_key(), algorithm="RS256", headers={"kid": "fake-ci"})


# A fake Google Drive (slice F2): shared drives and files smoke sets through POST /google/files, the OAuth token endpoint,
# and the Drive v3 calls the connector makes, so smoke can prove sync and permissions without Google.
DRIVES: dict[str, dict] = {}  # drive id -> {"name": ..., "files": {file id: the file, with its text}}
DRIVE_REFRESH, DRIVE_TOKEN = "smoke-refresh", "fake-drive-token"
# A shared drive's changes feed (H6): kept by comparing each POST /google/files with the last ({drive: {"seq", "changed":
# {file id: seq}, "gone": {file id: seq}, "last": {file id: file}}}); GET /google/changes-calls lists the feed's reads.
# POST /google/expire {"drive"} makes Drive refuse every token handed out so far, as it does a stale one (H6b).
DRIVE_CHANGES: dict = {}
DRIVE_CALLS: list = []
GOOGLE = os.environ.get("KL_FAKE_GOOGLE", "http://fake-llm:8080/google")  # the issuer of its sign-in ID tokens (slice F2b)
SIGNIN_CODES: dict[str, dict] = {}  # code -> who signed in; smoke sets these through POST /google/signin
# A fake Jira Cloud (ADR 0011): projects, issues, and who may see each, which smoke sets through POST /jira/setup, and the
# REST calls the connector makes, as its service account (an email and API token), so smoke proves the permission mapping.
JIRA: dict = {"projects": {}}
JIRA_AUTH = "jira-bot@acme.test:smoke-jira-token"
# A second site at /jira-b, set through POST /jira-b/setup: another company's Jira, whose project and issue keys may be
# the same as the first's, as two real sites' can be.
JIRA_SITES = {"jira": JIRA, "jira-b": {"projects": {}}, "retail-jira": {"projects": {}}}  # retail-jira: the synthetic retailer's (F9)
# A fake Confluence Cloud (S1): spaces, their pages and blog posts, and who may read each (the space's grants, each page's
# restrictions, and groups' members), which smoke sets through POST /confluence/setup. It answers the calls the
# connector makes, at Confluence's own paths (/confluence/wiki/...).
CONFLUENCE: dict = {"spaces": {}, "groups": {}, "users": {}}
# The synthetic retailer (F9) has its own site, at /retail-confluence, set through POST /retail-confluence/setup, so loading
# it never replaces the migration scenario's content (each setup replaces its site's whole state).
CONFLUENCE_SITES = {"confluence": CONFLUENCE, "retail-confluence": {"spaces": {}, "groups": {}, "users": {}}}
CONFLUENCE_AUTH = "confluence-bot@acme.test:smoke-confluence-token"
# Service accounts and shared folders (ADR 0011): smoke mints an account's key through POST /google/service-account (its
# token_uri points somewhere unreachable, so a worker that followed it would fail) and sets folders through POST
# /google/folders, each shared with the account or not.
SA_KEYS: dict = {}  # an account's email -> the public half of its key
FOLDERS: dict[str, dict] = {}  # folder id -> {"name", "shared": with the service account, "files": [files and subfolders]}
FOLDER_TYPE = "application/vnd.google-apps.folder"
# A fake Okta (ADR 0011: directories): users smoke sets through POST /okta/setup. Like Okta, its user list leaves the
# deprovisioned out unless they're searched for by status, so a directory has to ask for leavers to find them.
OKTA: dict = {"users": []}
OKTA_TOKEN = "smoke-okta-token"
# A fake Google Workspace Directory API (S2): users, recently deleted users, groups, and their members, which smoke sets
# through POST /google/admin/setup. Its token comes from the service-account token endpoint below, only for the three
# read-only directory scopes and an admin it was told to allow (domain-wide delegation's `sub`). Every token request is
# kept, so smoke can see what was asked for and where it went (GET /google/admin/requests). Pages are small on purpose.
ADMIN: dict = {"users": [], "deleted": [], "groups": {}, "subjects": [], "broken": False}
ADMIN_TOKEN = "fake-admin-token"
ADMIN_SCOPES = ["https://www.googleapis.com/auth/admin.directory.user.readonly",
                "https://www.googleapis.com/auth/admin.directory.group.readonly",
                "https://www.googleapis.com/auth/admin.directory.group.member.readonly"]
TOKEN_REQUESTS: list = []
ADMIN_PAGE = 2
# A fake Microsoft Graph and Entra token endpoint (S3): smoke mints an app registration with its own certificate
# (POST /microsoft/app), sets users, deleted users, and groups (POST /microsoft/setup), and the client signs in with a
# client assertion this fake verifies as Entra does: PS256, the certificate's SHA-256 thumbprint, this token address
# as its audience, the app as issuer and subject. Pages are small and come with @odata.nextLink, as Graph's do.
MICROSOFT = os.environ.get("KL_FAKE_MICROSOFT", "http://fake-llm:8080/microsoft")
GRAPH: dict = {"apps": {}, "users": [], "deleted": [], "groups": {}, "next_elsewhere": False, "sites": {}, "teams": {}}
# Teams (C1a): teams, each consented to by the apps its owner installed (resource-specific consent), with members and
# channels holding posts and their replies in Graph's own shapes; set through POST /microsoft/teams. A team an app has
# no consent for answers 403, as Graph does.
# Bot Framework (C1a): its OpenID metadata and keys (one endorsed for Teams, one only for Skype), a signer smoke uses
# to make the tokens Bot Framework sends (POST /botframework/sign), and a Bot Connector that takes the bot's posts for
# its app's token and keeps them (GET /teams-bot/posted).
BOTFRAMEWORK = os.environ.get("KL_FAKE_BOTFRAMEWORK", "http://fake-llm:8080/botframework")
BOT_SCOPE = "https://api.botframework.com/.default"
BOT_TOKEN = "fake-bot-token"
TEAMS_BOT: dict = {"posted": []}
_BOT_KEYS: dict = {}
# SharePoint, for Microsoft 365 (S5): sites, each granted to the apps an admin chose (Sites.Selected), with document
# libraries holding folders and files, each file with its permissions and text; set through POST /microsoft/sites. A
# file's content is a redirect to a download address, which smoke can point at another host (bad_download).
GRAPH_TOKEN = "fake-graph-token"
# Graph's delta (H6): each library's changes, kept by comparing each POST /microsoft/sites with the last ({library:
# {"seq", "changed": {item id: seq}, "gone": {item id: seq}, "last": {item id: item}}}); a delta with ?token=N answers
# what changed after N, deletions as {"id", "deleted": {}}. GET /microsoft/deltas lists the delta calls made.
# POST /microsoft/expire {"drive"} expires every delta link handed out so far: Graph's 410 resyncRequired (H6b).
DELTA: dict = {}
DELTA_CALLS: list = []
GRAPH_REQUESTS: list = []
# A fake Slack Web API (S4): a workspace, its people, and channels with members and messages (threads with their
# replies), which smoke sets through POST /slack/setup. Like Slack it answers 200 with ok: false for a bad token, pages
# with an opaque cursor, and, when smoke asks (throttle), answers one history call with 429 and Retry-After.
SLACK: dict = {"team": {"id": "T0SMOKE", "name": "Acme", "url": "https://acme-smoke.slack.com"}, "users": [], "channels": {},
               "throttle": False}
# The synthetic retailer's own workspace (F9), at /retail-slack, as the retailer's Confluence site is.
SLACK_SITES = {"slack": SLACK, "retail-slack": {"team": {"id": "T0RETAIL", "name": "Acme Outfitters", "url": "https://acme-outfitters.slack.com"},
                                                "users": [], "channels": {}, "throttle": False}}
SLACK_BOT = "UAERUGOBOT"  # the bot's own user id, as auth.test gives it (C2d: a mention names it)
SLACK_TOKEN = "xoxb-smoke-slack-token"
SLACK_PAGE = 2
# A fake Notion (S6): the pages and data sources shared with an integration (and some that aren't), with their blocks,
# which smoke sets through POST /notion/setup. It insists on the token and the Notion-Version header, as Notion does,
# and pages search, queries, and block children with cursors.
NOTION: dict = {"workspace": "Acme Notion", "pages": {}, "sources": {}}
NOTION_TOKEN = "ntn_smoke_notion_integration_token"
NOTION_VERSION = "2026-03-11"
# A fake GCP (ADR 0011): the metadata server's token, for the runner's own workload identity, and Secret Manager's
# access endpoint, holding the secrets smoke sets through POST /gcp/secrets as JSON objects.
GCP_SECRETS: dict = {}
GCP_TOKEN = "fake-gcp-workload-token"
# Fake data catalogs (S7): one set of tables, which smoke sets through POST /catalog/setup, served in each catalog's own
# shapes and behind each one's own sign-in: Dataplex with BigQuery's access and the Data Lineage API (the runner's GCP
# token); Unity Catalog (OAuth machine-to-machine); Glue and Lake Formation (STS trading the runner's identity token
# for a role, and every call SigV4-signed); Purview (an Entra app's certificate, for Purview's audience); and DataHub's
# GraphQL (an access token). A table: {"name": "<db>.<table>", "description", "columns": [[name, type, description]],
# "owner": email, "users": [emails who may read it], "groups": [group emails, granting through a directory only],
# "domain": bool (the company's domain may read it), "upstream": ["<db>.<table>"]}; "people": {email: Entra object id},
# for Purview, which names people by id. Pages hold one table, so every client's paging is exercised.
CATALOG: dict = {"tables": [], "people": {}}
# DataHub's own platforms beside CATALOG's Snowflake (F10b: the synthetic enterprise's EDW, on Teradata), set through
# POST /datahub/setup, so loading the enterprise never replaces smoke's tables, nor smoke the enterprise's.
DATAHUB_PLATFORMS: dict = {}
# Fake secret stores (G3): HashiCorp Vault (JWT login with the runner's identity token, then KV v2), AWS Secrets
# Manager (behind the same STS and SigV4 as Glue), and Azure Key Vault (an Entra token by workload identity
# federation), each holding the secrets smoke sets through POST /secret-stores as JSON strings.
SECRET_STORES: dict = {"vault": {}, "aws": {}, "azure": {}}
VAULT_TOKEN = "fake-vault-client-token"
AZURE_KV_TOKEN = "fake-azure-key-vault-token"
CATALOG_PROJECT = "acme-data"
DATABRICKS = {"client_id": "0b1e6a4e-6d0c-4c7e-9a43-5f2d8c1e7a10", "client_secret": "dbx-smoke-oauth-secret", "token": "dbx-smoke-token"}
AWS = {"role": "arn:aws:iam::123456789012:role/aerugo-catalog", "key_id": "ASIAFAKESMOKE", "secret": "fake-aws-secret-key",
       "session": "fake-aws-session-token", "saml": "arn:aws:iam::123456789012:saml-provider/acme"}
GCP_IDENTITY = "fake-gcp-identity-token-for-sts.amazonaws.com"
PURVIEW = os.environ.get("KL_FAKE_PURVIEW", "http://fake-llm:8080/purview")
PURVIEW_TOKEN = "fake-purview-token"
DATAHUB_TOKEN = "datahub-smoke-access-token-0123456789"


def new_service_account(email: str) -> dict:
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    SA_KEYS[email] = key.public_key()
    pem = key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()).decode()
    return {"type": "service_account", "client_email": email, "private_key": pem, "token_uri": "http://10.255.255.1/steal"}


def service_account_token(assertion: str) -> tuple[int, dict]:
    """Google's token endpoint for a service account: a JWT the account signed, aimed at this endpoint, for read-only
    Drive, or (S2) for exactly the three read-only directory scopes acting as an admin delegation allows."""
    import jwt as pyjwt
    try:
        iss = pyjwt.decode(assertion, options={"verify_signature": False}).get("iss")
        claims = pyjwt.decode(assertion, SA_KEYS[iss], algorithms=["RS256"], audience=f"{GOOGLE}/token")
    except Exception:  # noqa: BLE001 - any bad assertion is a refused grant
        return 400, {"error": "invalid_grant"}
    TOKEN_REQUESTS.append({"iss": claims.get("iss"), "scope": claims.get("scope"), "sub": claims.get("sub"), "aud": claims.get("aud")})
    if claims.get("scope") == "https://www.googleapis.com/auth/drive.readonly" and not claims.get("sub"):
        return 200, {"access_token": DRIVE_TOKEN, "token_type": "Bearer", "expires_in": 3600}
    if sorted(str(claims.get("scope") or "").split(" ")) == sorted(ADMIN_SCOPES):
        if claims.get("sub") not in ADMIN["subjects"]:
            return 401, {"error": "unauthorized_client", "error_description": "Client is unauthorized to retrieve access tokens"}
        return 200, {"access_token": ADMIN_TOKEN, "token_type": "Bearer", "expires_in": 3600}
    return 400, {"error": "invalid_scope"}


def new_entra_app(tenant: str, client: str | None = None) -> dict:
    """An app registration and its certificate: the PEM bundle the client takes, and the public half Entra keeps. Given
    an id, the app is that one, with a new certificate (the retailer's Teams app keeps its id across the fake's restarts)."""
    import base64, datetime, hashlib, uuid as uuidlib
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.x509.oid import NameOID
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "aerugo-directory")])
    now = datetime.datetime.now(datetime.timezone.utc)
    cert = (x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key())
            .serial_number(x509.random_serial_number()).not_valid_before(now - datetime.timedelta(minutes=5))
            .not_valid_after(now + datetime.timedelta(days=1)).sign(key, hashes.SHA256()))
    client = client or str(uuidlib.uuid4())
    der = cert.public_bytes(serialization.Encoding.DER)
    GRAPH["apps"][client] = {"tenant": tenant.lower(), "key": key.public_key(),
                             "thumbprint": base64.urlsafe_b64encode(hashlib.sha256(der).digest()).rstrip(b"=").decode()}
    bundle = key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()).decode() \
        + cert.public_bytes(serialization.Encoding.PEM).decode()
    return {"tenant_id": tenant.lower(), "client_id": client, "certificate": bundle}


def entra_token(tenant: str, form: dict, login: str = MICROSOFT, scope: str = "https://graph.microsoft.com/.default",
                token: str = GRAPH_TOKEN) -> tuple[int, dict]:
    """Entra's token endpoint, the client credentials grant with a certificate, checked as Entra checks it: for Graph,
    or (S7) for Purview's own audience at its fake's address."""
    import jwt as pyjwt
    field = lambda k: form.get(k, [""])[0]
    if field("grant_type") != "client_credentials" or field("client_assertion_type") != "urn:ietf:params:oauth:client-assertion-type:jwt-bearer":
        return 400, {"error": "unsupported_grant_type"}
    if field("scope") != scope:
        return 400, {"error": "invalid_scope"}
    app = GRAPH["apps"].get(field("client_id"))
    assertion = field("client_assertion")
    try:
        header = pyjwt.get_unverified_header(assertion)
        if not app or app["tenant"] != tenant or header.get("alg") != "PS256" or header.get("x5t#S256") != app["thumbprint"]:
            raise ValueError("unknown app, tenant, algorithm, or certificate")
        claims = pyjwt.decode(assertion, app["key"], algorithms=["PS256"], audience=f"{login}/{tenant}/oauth2/v2.0/token",
                              options={"require": ["exp", "nbf", "iat", "jti", "iss", "sub"]})
        if claims["iss"] != field("client_id") or claims["sub"] != field("client_id") or claims["exp"] - claims["nbf"] > 600:
            raise ValueError("the assertion isn't the app's own, or lives too long")
    except Exception as e:  # noqa: BLE001 - any bad assertion is a refused client
        return 401, {"error": "invalid_client", "error_description": str(e)[:200]}
    GRAPH_REQUESTS.append({"tenant": tenant, "client_id": field("client_id"), "aud": claims["aud"], "alg": header["alg"],
                           "x5t#S256": header["x5t#S256"], "scope": field("scope")})
    return 200, {"access_token": f"{token}.{field('client_id')}", "token_type": "Bearer", "expires_in": 3599}  # the app's own


def bot_key(kid: str):
    """One of Bot Framework's signing keys: fake-bot is endorsed for Teams, fake-skype only for Skype."""
    if kid not in _BOT_KEYS:
        from cryptography.hazmat.primitives.asymmetric import rsa
        _BOT_KEYS[kid] = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    return _BOT_KEYS[kid]


def bot_jwks() -> dict:
    import jwt
    keys = []
    for kid, endorsed in (("fake-bot", ["msteams", "skype"]), ("fake-skype", ["skype"])):
        keys.append({**json.loads(jwt.algorithms.RSAAlgorithm.to_jwk(bot_key(kid).public_key())), "kid": kid, "use": "sig",
                     "endorsements": endorsed})
    return {"keys": keys}


def bot_token(request: dict) -> str:
    """A token as Bot Framework sends one to a bot: for an app, naming the service address the activity came from."""
    import jwt
    now = int(time.time()) + int(request.get("shift", 0))
    claims = {"iss": request.get("iss", "https://api.botframework.com"), "aud": request.get("aud"), "serviceurl": request.get("serviceurl"),
              "nbf": now - 60, "exp": now + int(request.get("ttl", 3600))}
    return jwt.encode(claims, bot_key(request.get("kid", "fake-bot")), algorithm="RS256", headers={"kid": request.get("kid", "fake-bot")})


def graph_transitive(group: str, seen: frozenset = frozenset()) -> list[dict]:
    """A group's users, nested groups flattened, as Graph's transitiveMembers lists them."""
    users = []
    for m in (GRAPH["groups"].get(group) or {}).get("members", []):
        if m.get("group") and m["group"] not in seen:
            users += graph_transitive(m["group"], seen | {group})
        elif m.get("user"):
            users += [u for u in GRAPH["users"] if u["id"] == m["user"]]
    return users


def admin_members(group: str, derived: bool, seen: frozenset = frozenset()) -> list[dict]:
    """A group's members as the Directory API lists them; with derived membership, nested groups' members too."""
    out = []
    for m in ADMIN["groups"].get(group, []):
        out.append(m)
        if derived and m.get("type") == "GROUP" and m["email"] not in seen:
            out += admin_members(m["email"], True, seen | {group})
    return out


def google_id_token(client_id: str, who: dict) -> str:
    """An ID token as Google's sign-in issues one: the account's email, and whether Google verified it."""
    import jwt
    now = int(time.time())
    claims = {"iss": GOOGLE, "aud": who.get("aud") or client_id, "sub": f"fake-{who['email']}", "email": who["email"],
              "email_verified": who.get("verified", True), "iat": now, "exp": now + 300}
    return jwt.encode(claims, ci_key(), algorithm="RS256", headers={"kid": "fake-ci"})


def drive_export(spec: dict, mime: str) -> bytes:
    """A Google Doc as Word, or a Sheet as Excel, from the file's text (the first line is the title or header)."""
    lines = spec.get("text", "").split("\n")
    buf = io.BytesIO()
    if mime.endswith("wordprocessingml.document"):
        from docx import Document
        doc = Document()
        doc.add_paragraph(lines[0], style="Title")
        for line in lines[1:]:
            doc.add_paragraph(line)
        doc.save(buf)
        return buf.getvalue()
    if mime.endswith("spreadsheetml.sheet"):
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        rows = "".join(f'<row r="{r}">' + "".join(f'<c r="{chr(65 + c)}{r}" t="inlineStr"><is><t>{v}</t></is></c>'
                                                 for c, v in enumerate(line.split(","))) + "</row>" for r, line in enumerate(lines, 1))
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("xl/workbook.xml", f'<workbook xmlns="{ns}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                       '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>')
            z.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                       '<Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
            z.writestr("xl/worksheets/sheet1.xml", f'<worksheet xmlns="{ns}"><sheetData>{rows}</sheetData></worksheet>')
        return buf.getvalue()
    return spec.get("text", "").encode()


def words(text: str) -> set[str]:
    return {w for w in re.findall(r"[a-z0-9][a-z0-9-]*", text.lower()) if len(w) > 3 and w not in STOP}


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def _json(self, status: int, body) -> None:
        data = json.dumps(body).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _body(self) -> dict:
        return json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0))) or b"{}")

    def _throttled(self) -> bool:
        """Answer this call as throttled, when smoke asked for that (H2)."""
        with THROTTLE_LOCK:
            if THROTTLE["left"] <= 0 or not THROTTLE["prefix"] or not self.path.startswith(THROTTLE["prefix"]):
                return False
            THROTTLE["left"] -= 1
            THROTTLE["answered"] += 1
            status, after = THROTTLE["status"], THROTTLE["retry_after"]
        self.rfile.read(int(self.headers.get("Content-Length", 0) or 0))
        self.send_response(status)
        if after is not None:
            self.send_header("Retry-After", str(after))
        self.send_header("Content-Length", "0")
        self.end_headers()
        return True

    def do_GET(self):
        if self.path == "/confluence/listings":  # each page and blog post listing, and its order (H6)
            return self._json(200, CONFLUENCE.get("listings", []))
        if self.path == "/slack/history-calls":  # each channel history read, and from when (H6)
            return self._json(200, SLACK.get("history_calls", []))
        if self.path == "/google/changes-calls":  # the changes feed's reads, and from which token (H6)
            return self._json(200, DRIVE_CALLS)
        if self.path == "/microsoft/deltas":  # the delta calls made, and from which token (H6)
            return self._json(200, DELTA_CALLS)
        if self.path == "/throttle":
            return self._json(200, {"answered": THROTTLE["answered"], "left": THROTTLE["left"]})
        if self._throttled():
            return None
        if self.path == "/requests":
            return self._json(200, REQUESTS)
        if self.path == "/reranks":
            return self._json(200, RERANKS)
        if self.path == "/alerts":
            return self._json(200, ALERTS)
        if self.path == "/oidc/.well-known/openid-configuration":
            return self._json(200, {"issuer": ISSUER, "jwks_uri": f"{ISSUER}/.well-known/jwks"})
        if self.path == "/oidc/.well-known/jwks":
            return self._json(200, ci_jwks())
        if self.path.startswith("/google/drive/v3/"):
            return self._drive(urlparse(self.path))
        if self.path == "/google/admin/requests":
            return self._json(200, TOKEN_REQUESTS)
        if self.path == "/microsoft/requests":
            return self._json(200, GRAPH_REQUESTS)
        if self.path == "/botframework/v1/.well-known/openidconfiguration":
            return self._json(200, {"issuer": "https://api.botframework.com", "jwks_uri": f"{BOTFRAMEWORK}/v1/.well-known/keys"})
        if self.path == "/botframework/v1/.well-known/keys":
            return self._json(200, bot_jwks())
        if self.path == "/teams-bot/posted":  # what the bot posted, for smoke to read (C1a)
            return self._json(200, TEAMS_BOT["posted"])
        if self.path.startswith("/microsoft/v1.0/"):
            return self._graph(urlparse(self.path))
        if self.path.startswith("/microsoft/download/"):  # a pre-authenticated download address: no Graph token
            _, _, _, drive, item = self.path.split("/", 4)
            spec = next((i for s in GRAPH["sites"].values() for d, lib in s["drives"].items() if d == drive
                         for i in lib["items"] if i["id"] == item), None)
            if spec is None:
                return self._json(404, {})
            data = spec.get("text", "").encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/octet-stream")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
            return
        site = self.path.split("/")[1]
        if site in SLACK_SITES and self.path.startswith(f"/{site}/api/"):
            return self._slack(urlparse(self.path), SLACK_SITES[site])
        if site in SLACK_SITES and self.path == f"/{site}/posted":  # what the bot posted, for smoke to read (C2b)
            return self._json(200, SLACK_SITES[site].get("posted", []))
        if self.path.startswith("/notion/v1/"):
            return self._notion("GET", urlparse(self.path), {})
        if self.path.startswith("/google/admin/directory/v1/"):
            return self._admin(urlparse(self.path))
        if self.path == "/google/oauth2/v3/certs":  # the keys Google signs ID tokens with
            return self._json(200, ci_jwks())
        if site in JIRA_SITES and self.path == f"/{site}/searches":  # the JQL each search asked for since its setup (H6)
            return self._json(200, JIRA_SITES[site].get("searches", []))
        if site in JIRA_SITES and self.path.startswith(f"/{site}/rest/api/3/"):
            return self._jira(urlparse(self.path), JIRA_SITES[site])
        if site in CONFLUENCE_SITES and self.path.startswith(f"/{site}/wiki/download/attachments/"):  # an attachment's bytes (Q7)
            import base64
            if self.headers.get("Authorization") != "Basic " + base64.b64encode(CONFLUENCE_AUTH.encode()).decode():
                return self._json(401, {"message": "Unauthorized"})
            aid = urlparse(self.path).path.rsplit("/", 1)[-1]
            found = next((a for s in CONFLUENCE_SITES[site]["spaces"].values() for kind in ("pages", "blogposts") for c in s.get(kind, [])
                          for a in c.get("attachments", []) if isinstance(a, dict) and a.get("id") == aid), None)
            return self._bytes(found["text"].encode()) if found else self._json(404, {})
        if site in CONFLUENCE_SITES and self.path.startswith(f"/{site}/wiki/"):
            return self._confluence(urlparse(self.path), site)
        if self.path.startswith("/okta/api/v1/users"):
            if self.headers.get("Authorization") != f"SSWS {OKTA_TOKEN}":
                return self._json(401, {"errorCode": "E0000011", "errorSummary": "Invalid token provided"})
            search = {k: v[0] for k, v in parse_qs(urlparse(self.path).query).items()}.get("search", "")
            gone = 'status eq "DEPROVISIONED"' in search
            return self._json(200, [u for u in OKTA["users"] if (u.get("status") == "DEPROVISIONED") == gone])
        if urlparse(self.path).path == "/gcp-metadata/computeMetadata/v1/instance/service-accounts/default/identity":
            if self.headers.get("Metadata-Flavor") != "Google":
                return self._json(403, {"error": "Missing Metadata-Flavor: Google"})
            audience = {k: v[0] for k, v in parse_qs(urlparse(self.path).query).items()}.get("audience")
            return self._xml(200, GCP_IDENTITY if audience == "sts.amazonaws.com" else f"fake-gcp-identity-token-for-{audience}") \
                if audience else self._json(400, {"error": "audience"})
        if self.path.startswith(("/gcp-catalog/", "/databricks/", "/purview/")):
            return self._catalog("GET", urlparse(self.path), b"")
        if self.path.startswith("/hcvault/v1/"):  # Vault's KV v2, for its login's token only
            if self.headers.get("X-Vault-Token") != VAULT_TOKEN:
                return self._json(403, {"errors": ["permission denied"]})
            key = urlparse(self.path).path.split("/data/", 1)[-1]
            value = SECRET_STORES["vault"].get(key)
            return self._json(200, {"data": {"data": json.loads(value)}}) if value else self._json(404, {"errors": []})
        if self.path.startswith("/azkv/"):  # Key Vault: /azkv/<vault>/secrets/<name>
            if self.headers.get("Authorization") != f"Bearer {AZURE_KV_TOKEN}":
                return self._json(401, {"error": {"code": "Unauthorized"}})
            parts = urlparse(self.path).path.split("/")
            value = SECRET_STORES["azure"].get(f"{parts[2]}/{parts[4]}") if len(parts) == 5 and parts[3] == "secrets" else None
            return self._json(200, {"value": value, "id": self.path}) if value else self._json(404, {"error": {"code": "SecretNotFound"}})
        if self.path == "/gcp-metadata/computeMetadata/v1/instance/service-accounts/default/token":
            if self.headers.get("Metadata-Flavor") != "Google":  # as the real metadata server insists
                return self._json(403, {"error": "Missing Metadata-Flavor: Google"})
            return self._json(200, {"access_token": GCP_TOKEN, "expires_in": 3599, "token_type": "Bearer"})
        if self.path.startswith("/gcp/v1/projects/") and self.path.endswith(":access"):
            if self.headers.get("Authorization") != f"Bearer {GCP_TOKEN}":
                return self._json(401, {"error": {"code": 401, "status": "UNAUTHENTICATED"}})
            import base64
            parts = self.path.split("/")  # /gcp/v1/projects/<p>/secrets/<s>/versions/<v>:access
            value = GCP_SECRETS.get(parts[6]) if len(parts) == 9 and parts[3] == "projects" and parts[5] == "secrets" else None
            if value is None:
                return self._json(404, {"error": {"code": 404, "status": "NOT_FOUND"}})
            return self._json(200, {"name": self.path[len("/gcp/v1/"):-len(":access")],
                                    "payload": {"data": base64.b64encode(json.dumps(value).encode()).decode()}})
        self._json(404, {})

    def _catalog(self, method: str, url, raw: bytes):
        """Every fake catalog's API (S7), each behind its own sign-in, all serving CATALOG's tables."""
        import base64, uuid as uuidlib
        path, q, tables = url.path, {k: v[0] for k, v in parse_qs(url.query).items()}, CATALOG["tables"]
        parts = lambda t: t["name"].split(".", 1)
        page = lambda items, start: (items[start:start + 1], str(start + 1) if start + 1 < len(items) else None)
        body = json.loads(raw or b"{}") if raw[:1] in (b"{", b"") else {}
        if path.startswith("/gcp-catalog/"):  # Dataplex, BigQuery's access, and Data Lineage, as the runner's own identity
            if self.headers.get("Authorization") != f"Bearer {GCP_TOKEN}":
                return self._json(401, {"error": {"code": 401, "status": "UNAUTHENTICATED"}})
            entry = lambda t: f"projects/{CATALOG_PROJECT}/locations/us/entryGroups/@bigquery/entries/{t['name'].replace('.', '-')}"
            fqn = lambda name: f"bigquery:{CATALOG_PROJECT}.{name}"
            if method == "POST" and path == f"/gcp-catalog/dataplex/v1/projects/{CATALOG_PROJECT}/locations/global:searchEntries":
                if q.get("query") != f"projectid={CATALOG_PROJECT}":
                    return self._json(400, {"error": {"status": "INVALID_ARGUMENT"}})
                found, following = page(tables, int(q.get("pageToken") or 0))
                return self._json(200, {"results": [{"dataplexEntry": {"name": entry(t)}} for t in found],
                                        **({"nextPageToken": following} if following else {})})
            if method == "POST" and path.startswith(f"/gcp-catalog/datalineage/v1/projects/{CATALOG_PROJECT}/locations/") and path.endswith(":searchLinks"):
                target = (body.get("target") or {}).get("fullyQualifiedName")
                return self._json(200, {"links": [{"source": {"fullyQualifiedName": fqn(up)}, "target": {"fullyQualifiedName": target}}
                                                  for t in tables if fqn(t["name"]) == target for up in t.get("upstream", [])]})
            for t in tables:
                ds, tb = parts(t)
                if method == "GET" and path == f"/gcp-catalog/dataplex/v1/{entry(t)}" and q.get("view") == "ALL":
                    return self._json(200, {
                        "name": entry(t), "fullyQualifiedName": fqn(t["name"]), "updateTime": t.get("version", "1"),
                        "entryType": "projects/655216118709/locations/global/entryTypes/bigquery-table",
                        "entrySource": {"system": "BigQuery", "description": t.get("description", ""), "labels": {"tier": "gold"},
                                        "resource": f"//bigquery.googleapis.com/projects/{CATALOG_PROJECT}/datasets/{ds}/tables/{tb}"},
                        "aspects": {"655216118709.global.schema": {"data": {"fields": [{"name": n, "dataType": ty, "description": d}
                                                                                       for n, ty, d in t.get("columns", [])]}},
                                    "655216118709.global.contacts": {"data": {"identities": [{"role": "Owner", "id": t["owner"]}] if t.get("owner") else []}}}})
                if method == "GET" and path == f"/gcp-catalog/bigquery/bigquery/v2/projects/{CATALOG_PROJECT}/datasets/{ds}":
                    access = [{"role": "READER", "specialGroup": "projectReaders"}, {"role": "READER", "iamMember": "allUsers"}]
                    access += [{"role": "READER", "domain": "acme.test"} for o in tables if parts(o)[0] == ds and o.get("domain")]
                    return self._json(200, {"access": access})
                if method == "POST" and path == f"/gcp-catalog/bigquery/bigquery/v2/projects/{CATALOG_PROJECT}/datasets/{ds}/tables/{tb}:getIamPolicy":
                    return self._json(200, {"bindings": [
                        {"role": "roles/bigquery.dataViewer", "members": [f"user:{u}" for u in t.get("users", [])] + [f"group:{g}" for g in t.get("groups", [])]
                         + ["allAuthenticatedUsers", "serviceAccount:etl@acme-data.iam.gserviceaccount.com"]},
                        {"role": "roles/bigquery.metadataViewer", "members": ["user:mallory@acme.test"]}]})
            return self._json(404, {"error": {"code": 404, "status": "NOT_FOUND"}})
        if path.startswith("/databricks/"):  # Unity Catalog, as a service principal
            if path == "/databricks/oidc/v1/token" and method == "POST":
                form = parse_qs(raw.decode())
                basic = base64.b64encode(f"{DATABRICKS['client_id']}:{DATABRICKS['client_secret']}".encode()).decode()
                if self.headers.get("Authorization") != f"Basic {basic}" or form.get("grant_type") != ["client_credentials"]:
                    return self._json(401, {"error": "invalid_client"})
                return self._json(200, {"access_token": DATABRICKS["token"], "token_type": "Bearer", "expires_in": 3600})
            if self.headers.get("Authorization") != f"Bearer {DATABRICKS['token']}":
                return self._json(401, {"error_code": "UNAUTHENTICATED"})
            rest = path[len("/databricks/api/"):]
            if rest == "2.1/unity-catalog/catalogs":
                return self._json(200, {"catalogs": [{"name": "main"}, {"name": "bad name!"}]})
            if rest == "2.1/unity-catalog/schemas":  # every table is in main; any other catalog is empty
                found, following = page(sorted({parts(t)[0] for t in tables}) if q.get("catalog_name") == "main" else [], int(q.get("page_token") or 0))
                return self._json(200, {"schemas": [{"name": s} for s in found], **({"next_page_token": following} if following else {})})
            if rest == "2.1/unity-catalog/tables" and q.get("catalog_name") == "main":
                mine = [t for t in tables if parts(t)[0] == q.get("schema_name")]
                found, following = page(mine, int(q.get("page_token") or 0))
                return self._json(200, {"tables": [{"full_name": f"main.{t['name']}", "table_type": "MANAGED", "comment": t.get("description", ""),
                                                    "owner": t.get("owner", ""), "updated_at": int(t.get("version", 1)),
                                                    "columns": [{"name": n, "type_text": ty, "comment": d} for n, ty, d in t.get("columns", [])]}
                                                   for t in found], **({"next_page_token": following} if following else {})})
            for t in tables:
                if rest == f"2.1/unity-catalog/effective-permissions/table/main.{t['name']}":
                    return self._json(200, {"privilege_assignments": [
                        {"principal": u, "privileges": [{"privilege": "SELECT", "inherited_from_type": "SCHEMA"}]} for u in t.get("users", [])]
                        + [{"principal": g.split("@")[0], "privileges": [{"privilege": "SELECT"}]} for g in t.get("groups", [])]
                        + [{"principal": "mallory@acme.test", "privileges": [{"privilege": "BROWSE"}]}]})
            if rest == "2.0/lineage-tracking/table-lineage":
                return self._json(200, {"upstreams": [{"tableInfo": {"catalog_name": "main", "schema_name": up.split(".")[0], "name": up.split(".")[1]}}
                                                      for t in tables if f"main.{t['name']}" == q.get("table_name") for up in t.get("upstream", [])]})
            return self._json(404, {"error_code": "NOT_FOUND"})
        if path.startswith("/aws/"):  # STS, Glue, and Lake Formation
            if path.rstrip("/") == "/aws/sts":
                form = {k: v[0] for k, v in parse_qs(raw.decode()).items()}
                if form.get("Action") != "AssumeRoleWithWebIdentity" or form.get("WebIdentityToken") != GCP_IDENTITY or form.get("RoleArn") != AWS["role"]:
                    return self._xml(403, "<ErrorResponse><Error><Code>AccessDenied</Code></Error></ErrorResponse>")
                return self._xml(200, '<AssumeRoleWithWebIdentityResponse xmlns="https://sts.amazonaws.com/doc/2011-06-15/">'
                                      "<AssumeRoleWithWebIdentityResult><Credentials>"
                                      f"<AccessKeyId>{AWS['key_id']}</AccessKeyId><SecretAccessKey>{AWS['secret']}</SecretAccessKey>"
                                      f"<SessionToken>{AWS['session']}</SessionToken><Expiration>2099-01-01T00:00:00Z</Expiration>"
                                      "</Credentials></AssumeRoleWithWebIdentityResult></AssumeRoleWithWebIdentityResponse>")
            service = path.split("/")[2]
            if not self._signed_by_role(service, path, raw):
                return self._json(403, {"__type": "InvalidSignatureException"})
            if service == "secretsmanager":
                value = SECRET_STORES["aws"].get(body.get("SecretId"))
                if self.headers.get("X-Amz-Target") != "secretsmanager.GetSecretValue" or value is None:
                    return self._json(400, {"__type": "ResourceNotFoundException"})
                return self._json(200, {"Name": body["SecretId"], "SecretString": value})
            if service == "glue":
                target = self.headers.get("X-Amz-Target")
                if target == "AWSGlue.GetDatabases":
                    found, following = page(sorted({parts(t)[0] for t in tables}), int(body.get("NextToken") or 0))
                    return self._json(200, {"DatabaseList": [{"Name": d} for d in found], **({"NextToken": following} if following else {})})
                if target == "AWSGlue.GetTables":
                    mine = [t for t in tables if parts(t)[0] == body.get("DatabaseName")]
                    found, following = page(mine, int(body.get("NextToken") or 0))
                    return self._json(200, {"TableList": [{"Name": parts(t)[1], "DatabaseName": parts(t)[0], "Description": t.get("description", ""),
                                                           "Owner": t.get("owner", ""), "TableType": "EXTERNAL_TABLE", "UpdateTime": float(t.get("version", 1)),
                                                           "Parameters": {"classification": "parquet"},
                                                           "StorageDescriptor": {"Location": f"s3://acme-lake/{t['name']}/",
                                                                                 "Columns": [{"Name": n, "Type": ty, "Comment": d} for n, ty, d in t.get("columns", [])]}}
                                                          for t in found], **({"NextToken": following} if following else {})})
                return self._json(400, {"__type": "UnknownOperationException"})
            if service == "lakeformation" and path == "/aws/lakeformation/ListPermissions":
                table = (body.get("Resource") or {}).get("Table") or {}
                spec = next((t for t in tables if t["name"] == f"{table.get('DatabaseName')}.{table.get('Name')}"), None)
                if spec is None or body.get("IncludeRelated") != "TRUE":
                    return self._json(400, {"__type": "InvalidInputException"})
                res = {"Table": {"DatabaseName": table["DatabaseName"], "Name": table["Name"]}}
                grants = [{"Principal": {"DataLakePrincipalIdentifier": f"{AWS['saml']}:user/{u}"}, "Permissions": ["SELECT"], "Resource": res}
                          for u in spec.get("users", [])]
                grants += [{"Principal": {"DataLakePrincipalIdentifier": f"{AWS['saml']}:group/{g}"}, "Permissions": ["SELECT"], "Resource": res}
                           for g in spec.get("groups", [])]
                grants += [{"Principal": {"DataLakePrincipalIdentifier": "IAMAllowedPrincipals"}, "Permissions": ["ALL"], "Resource": res},
                           {"Principal": {"DataLakePrincipalIdentifier": f"{AWS['saml']}:user/mallory@acme.test"}, "Permissions": ["ALL"],
                            "Resource": {"Database": {"Name": table["DatabaseName"]}}},
                           {"Principal": {"DataLakePrincipalIdentifier": f"{AWS['saml']}:user/trent@acme.test"}, "Permissions": ["DESCRIBE"], "Resource": res}]
                return self._json(200, {"PrincipalResourcePermissions": grants})
            return self._json(404, {})
        if path.startswith("/purview/"):  # Purview's Data Map, as an Entra app with its certificate
            client = (self.headers.get("Authorization") or "").removeprefix(f"Bearer {PURVIEW_TOKEN}.")
            if client not in GRAPH["apps"]:
                return self._json(401, {"error": {"code": "Unauthorized"}})
            guid = lambda name: str(uuidlib.uuid5(uuidlib.NAMESPACE_URL, f"purview/{name}"))
            qualified = lambda name: f"mssql://acme.database.windows.net/{name.replace('.', '/dbo/')}"
            oid = lambda email: CATALOG["people"].get(email)
            if path == "/purview/policystore/metadataPolicies" and q.get("api-version") == "2021-07-01":
                values = []
                for c in sorted({parts(t)[0] for t in tables}):
                    readers = sorted({oid(u) for t in tables if parts(t)[0] == c for u in t.get("users", []) if oid(u)})
                    admins = sorted({oid(t["owner"]) for t in tables if parts(t)[0] == c and oid(t.get("owner", ""))})
                    rule = lambda role, ids: {"id": f"purviewmetadatarole_builtin_{role}:{c}", "dnfCondition": [[
                        {"attributeName": "principal.microsoft.id", "attributeValueIncludedIn": ids},
                        {"attributeName": "derived.purview.role", "attributeValueIncludedIn": [f"purviewmetadatarole_builtin_{role}"]}]]}
                    values.append({"name": f"policy_{c}", "properties": {"collection": {"referenceName": c, "type": "CollectionReference"},
                                                                         "attributeRules": [rule("data-reader", readers), rule("collection-administrator", admins)]}})
                return self._json(200, {"values": values})
            if path == "/purview/datamap/api/search/query" and method == "POST":
                mine = [t for t in tables if parts(t)[0] == (body.get("filter") or {}).get("collectionId")]
                found, following = page(mine, int(body.get("continuationToken") or 0))
                return self._json(200, {"value": [{"id": guid(t["name"]), "qualifiedName": qualified(t["name"]), "entityType": "azure_sql_table"} for t in found],
                                        **({"continuationToken": following} if following else {})})
            for t in tables:
                if path == f"/purview/datamap/api/atlas/v2/entity/guid/{guid(t['name'])}":
                    cols = {guid(f"{t['name']}#{n}"): {"typeName": "azure_sql_column", "attributes": {"name": n, "data_type": ty, "description": d}}
                            for n, ty, d in t.get("columns", [])}
                    return self._json(200, {"entity": {"guid": guid(t["name"]), "typeName": "azure_sql_table", "updateTime": int(t.get("version", 1)),
                                                       "attributes": {"qualifiedName": qualified(t["name"]), "name": parts(t)[1], "userDescription": t.get("description", "")},
                                                       "relationshipAttributes": {"columns": [{"guid": g} for g in cols]},
                                                       "contacts": {"Owner": [{"id": oid(t["owner"])}] if oid(t.get("owner", "")) else []}},
                                            "referredEntities": cols})
                if path == f"/purview/datamap/api/atlas/v2/lineage/{guid(t['name'])}" and q.get("direction") == "INPUT":
                    process = guid(f"{t['name']}#process")
                    nodes = {guid(t["name"]): {"typeName": "azure_sql_table", "attributes": {"qualifiedName": qualified(t["name"])}},
                             process: {"typeName": "azure_data_factory_process", "attributes": {"qualifiedName": f"adf://copy/{t['name']}"}}}
                    nodes.update({guid(up): {"typeName": "azure_sql_table", "attributes": {"qualifiedName": qualified(up)}} for up in t.get("upstream", [])})
                    return self._json(200, {"guidEntityMap": nodes})
            return self._json(404, {"error": {"code": "NotFound"}})
        if path == "/datahub/api/graphql" and method == "POST":  # DataHub's GraphQL, with an access token
            if self.headers.get("Authorization") != f"Bearer {DATAHUB_TOKEN}":
                return self._json(401, {"message": "Unauthorized"})
            if "scrollAcrossEntities" not in str(body.get("query")):
                return self._json(200, {"errors": [{"message": "unknown query"}]})
            given = (body.get("variables") or {}).get("input") or {}
            platform = ((given.get("orFilters") or [{"and": [{"values": ["urn:li:dataPlatform:snowflake"]}]}])[0]["and"][0]["values"][0])
            named = platform.rsplit(":", 1)[1]
            urn = lambda name: f"urn:li:dataset:(urn:li:dataPlatform:{named},{name},PROD)"
            mine = DATAHUB_PLATFORMS[named] if named in DATAHUB_PLATFORMS else tables if named == "snowflake" else []
            found, following = page(mine, int(given.get("scrollId") or 0))
            return self._json(200, {"data": {"scrollAcrossEntities": {"nextScrollId": following, "searchResults": [{"entity": {
                "urn": urn(t["name"]), "properties": {"description": t.get("description", "")}, "editableProperties": None,
                "schemaMetadata": {"fields": [{"fieldPath": n, "nativeDataType": ty, "description": d} for n, ty, d in t.get("columns", [])]},
                "ownership": {"owners": [{"owner": {"properties": {"email": u}}} for u in t.get("users", [])]
                              + [{"owner": {"name": g}} for g in t.get("groups", [])]},
                "tags": {"tags": [{"tag": {"properties": {"name": "gold"}}}]}, "glossaryTerms": {"terms": [{"term": {"properties": {"name": "Order"}}}]},
                "upstream": {"relationships": [{"entity": {"urn": urn(up)}} for up in t.get("upstream", [])]}}} for t in found]}}})
        return self._json(404, {})

    def _signed_by_role(self, service: str, path: str, raw: bytes) -> bool:
        """Whether an AWS call is SigV4-signed with the role's credentials, over exactly what was sent (S7)."""
        import datetime
        from .glue import Credentials, sign
        auth = self.headers.get("Authorization") or ""
        m = re.fullmatch(r"AWS4-HMAC-SHA256 Credential=([A-Z0-9]+)/(\d{8})/([a-z0-9-]+)/([a-z]+)/aws4_request, SignedHeaders=([a-z0-9;-]+), Signature=[0-9a-f]{64}", auth)
        if not m or m[1] != AWS["key_id"] or m[4] != service or self.headers.get("X-Amz-Security-Token") != AWS["session"]:
            return False
        names = m[5].split(";")
        given = {k: self.headers.get(k) or "" for k in names if k not in ("host", "x-amz-date", "x-amz-security-token")}
        when = datetime.datetime.strptime(self.headers.get("X-Amz-Date") or "", "%Y%m%dT%H%M%SZ")
        expected = sign("POST", f"http://{self.headers.get('Host')}{path}", m[3], service, raw, given,
                        Credentials(AWS["key_id"], AWS["secret"], AWS["session"]), when)
        return expected["authorization"] == auth

    def _bytes(self, data: bytes) -> None:
        self.send_response(200)
        self.send_header("Content-Type", "application/octet-stream")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _xml(self, status: int, body: str) -> None:
        data = body.encode()
        self.send_response(status)
        self.send_header("Content-Type", "text/xml")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def _drive(self, url):
        if self.headers.get("Authorization") != f"Bearer {DRIVE_TOKEN}":
            return self._json(401, {"error": {"message": "invalid credentials"}})
        q, rest = {k: v[0] for k, v in parse_qs(url.query).items()}, url.path[len("/google/drive/v3/"):]
        if rest == "drives":
            return self._json(200, {"drives": [{"id": d, "name": v["name"]} for d, v in DRIVES.items()]})
        if rest in ("changes/startPageToken", "changes"):  # a shared drive's changes feed (H6)
            state = DRIVE_CHANGES.get(q.get("driveId", ""), {"seq": 0, "changed": {}, "gone": {}})
            DRIVE_CALLS.append({"drive": q.get("driveId"), "call": rest, "token": q.get("pageToken")})
            if rest == "changes/startPageToken":
                return self._json(200, {"startPageToken": str(state["seq"])})
            after, files = int(q.get("pageToken") or 0), DRIVES.get(q.get("driveId", ""), {"files": {}})["files"]
            if after <= state.get("expired_to", -1):  # as Drive answers a token it no longer accepts (H6b)
                return self._json(400, {"error": {"code": 400, "message": "Invalid Value", "errors": [
                    {"domain": "global", "reason": "invalid", "message": "Invalid Value", "locationType": "parameter",
                     "location": "pageToken"}]}})
            changed = [{"fileId": fid, "removed": False, "file": {k: v for k, v in files[fid].items() if k != "text"}}
                       for fid, at in state["changed"].items() if at > after and fid in files]
            removed = [{"fileId": fid, "removed": True} for fid, at in state["gone"].items() if at > after]
            return self._json(200, {"changes": changed + removed, "newStartPageToken": str(state["seq"])})
        if rest == "files":
            query = q.get("q", "")
            if "sharedWithMe = true" in query:  # folders shared with the service account
                files = [{"id": fid, "name": f["name"], "mimeType": FOLDER_TYPE} for fid, f in FOLDERS.items() if f.get("shared")]
            elif "' in parents" in query:  # one folder's children: its files and its subfolders
                files = FOLDERS.get(query.split("'")[1], {"files": []})["files"]
            else:
                files = DRIVES.get(q.get("driveId", ""), {"files": {}})["files"].values()
            return self._json(200, {"files": [{k: v for k, v in f.items() if k != "text"} for f in files]})
        parts = rest.split("/")  # files/<id> or files/<id>/export
        spec = next((d["files"][parts[1]] for d in DRIVES.values() if len(parts) > 1 and parts[1] in d["files"]), None) or \
            next((f for d in FOLDERS.values() for f in d["files"] if len(parts) > 1 and f.get("id") == parts[1] and f.get("mimeType") != FOLDER_TYPE), None)
        if parts[0] != "files" or spec is None:
            return self._json(404, {"error": {"message": "file not found"}})
        body = drive_export(spec, q.get("mimeType", "")) if parts[-1] == "export" else spec.get("text", "").encode()
        self.send_response(200)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _confluence(self, url, site: str = "confluence"):
        import base64
        if self.headers.get("Authorization") != "Basic " + base64.b64encode(CONFLUENCE_AUTH.encode()).decode():
            return self._json(401, {"message": "Unauthorized"})
        params = parse_qs(url.query)
        q = {k: v[0] for k, v in params.items()}
        rest, state = url.path[len(f"/{site}/wiki/"):], CONFLUENCE_SITES[site]
        parts = rest.split("/")
        spaces, groups, users = state["spaces"], state["groups"], state["users"]
        by_id = {str(s["id"]): (k, s) for k, s in spaces.items()}
        content = {str(c["id"]): c for s in spaces.values() for kind in ("pages", "blogposts") for c in s.get(kind, [])}
        person = lambda a: {"accountId": a, "accountType": "atlassian", "email": (users.get(a) or {}).get("email"),
                            "displayName": (users.get(a) or {}).get("name") or a}
        read = lambda kind, who: {"principal": {"type": kind, "id": who}, "operation": {"key": "read", "targetType": "space"}}
        if rest == "rest/api/user/current":
            return self._json(200, {"accountId": "aerugo-bot", "displayName": "Aerugo bot"})
        kind_of = {str(c["id"]): kind for s in spaces.values() for kind in ("pages", "blogposts") for c in s.get(kind, [])}
        space_of = {str(c["id"]): str(s["id"]) for s in spaces.values() for kind in ("pages", "blogposts") for c in s.get(kind, [])}
        if len(parts) == 4 and parts[:3] in (["api", "v2", "pages"], ["api", "v2", "blogposts"]) and kind_of.get(parts[3]) == parts[2]:
            c = content[parts[3]]  # one page or post; its body only when asked for (a pointer's, read live: F11a)
            if q.get("body-format") == "storage":
                state.setdefault("page_reads", []).append(parts[3])
            return self._json(200, {"id": str(c["id"]), "title": c["title"], "parentId": c.get("parentId"), "spaceId": space_of[parts[3]],
                                    "version": {"number": c.get("version", 1), "createdAt": c.get("modified", "2026-01-01T00:00:00.000Z")},
                                    **({"body": {"storage": {"value": c.get("body", "")}}} if q.get("body-format") == "storage" else {})})
        if rest == "rest/api/search":  # CQL, as a live-only connection asks it (F11b): its spaces, any term
            cql = q.get("cql", "")
            state.setdefault("searches", []).append(cql)
            terms = re.findall(r'text ~ "([a-z0-9]+)"', cql)
            listed = re.search(r"space in \(([^)]*)\)", cql)
            wanted = {k.strip().strip('"') for k in listed[1].split(",")} if listed else set(spaces)
            words = lambda c: re.findall(r"[a-z0-9]+", re.sub(r"<[^>]+>", " ", f"{c['title']} {c.get('body', '')}").lower())
            hits = [(sum(t in words(c) for t in terms), n, c) for n, c in
                    enumerate(c for k in sorted(wanted) for kind in ("pages", "blogposts") for c in spaces.get(k, {}).get(kind, []))]
            found = [c for score, _, c in sorted(hits, key=lambda h: (-h[0], h[1])) if score]
            return self._json(200, {"results": [{"content": {"id": str(c["id"]), "type": "page", "title": c["title"]}}
                                                for c in found[:int(q.get("limit") or 25)]]})
        if rest == "api/v2/spaces":
            keys = set(q["keys"].split(",")) if q.get("keys") else None
            return self._json(200, {"results": [{"id": str(s["id"]), "key": k, "name": s["name"]} for k, s in spaces.items()
                                                if keys is None or k in keys], "_links": {}})
        if len(parts) == 5 and parts[:3] == ["api", "v2", "spaces"] and parts[3] in by_id:
            key, s = by_id[parts[3]]
            if parts[4] in ("pages", "blogposts"):  # newest-modified first when asked (H6), and each listing recorded
                state.setdefault("listings", []).append({"space": key, "kind": parts[4], "sort": q.get("sort")})
                listed = s.get(parts[4], [])
                if q.get("sort") == "-modified-date":
                    listed = sorted(listed, key=lambda c: c.get("modified", "2026-01-01T00:00:00.000Z"), reverse=True)
                return self._json(200, {"results": [
                    {"id": str(c["id"]), "title": c["title"], "parentId": c.get("parentId"),
                     "version": {"number": c.get("version", 1), "createdAt": c.get("modified", "2026-01-01T00:00:00.000Z")},
                     "body": {"storage": {"value": c.get("body", "")}}, "_links": {"webui": f"/spaces/{key}/pages/{c['id']}"}}
                    for c in listed], "_links": {}})
            if parts[4] == "permissions":
                grants = [read("user", a) for a in s.get("users", [])] + [read("group", g) for g in s.get("groups", [])]
                return self._json(200, {"results": grants + ([read("anonymous", None)] if s.get("anonymous") else []), "_links": {}})
        if len(parts) == 7 and parts[:3] == ["rest", "api", "content"] and parts[4:] == ["restriction", "byOperation", "read"]:
            c = content.get(parts[3])
            if c is None:
                return self._json(404, {"message": "No content found"})
            return self._json(200, {"operation": "read", "restrictions": {
                "user": {"results": [person(a) for a in c.get("restricted_users", [])]},
                "group": {"results": [{"id": g, "name": (groups.get(g) or {}).get("name", g)} for g in c.get("restricted_groups", [])]}}})
        if len(parts) == 5 and parts[:3] == ["rest", "api", "group"] and parts[4] == "membersByGroupId":
            if parts[3] not in groups:
                return self._json(404, {"message": "No group found"})
            return self._json(200, {"results": [person(a) for a in groups[parts[3]].get("members", [])]})
        if rest == "rest/api/group/by-id":
            g = groups.get(q.get("id", ""))
            return self._json(200, {"id": q.get("id"), "name": g["name"]}) if g else self._json(404, {"message": "No group found"})
        if rest == "rest/api/user/bulk":
            return self._json(200, {"results": [person(a) for a in params.get("accountId", []) if a in users]})
        if len(parts) == 5 and parts[:2] == ["api", "v2"] and parts[2] in ("pages", "blogposts") and parts[4] in ("labels", "attachments"):
            c = content.get(parts[3]) or {}
            if parts[4] == "labels":
                return self._json(200, {"results": [{"name": label} for label in c.get("labels", [])], "_links": {}})
            files = [a if isinstance(a, dict) else {"title": a} for a in c.get("attachments", [])]
            return self._json(200, {"results": [{"title": a["title"], **({"id": a["id"], "fileSize": len(a.get("text", "")),
                                                                          "version": {"number": 1}, "downloadLink": f"/download/attachments/{c['id']}/{a['id']}"}
                                                                         if a.get("id") else {})} for a in files], "_links": {}})
        return self._json(404, {"message": "Not found"})

    def _notion(self, method, url, body):
        """Notion's API, as the connector calls it: users/me, search, a data source's query, and block children."""
        if self.headers.get("Authorization") != f"Bearer {NOTION_TOKEN}":
            return self._json(401, {"object": "error", "code": "unauthorized", "message": "API token is invalid."})
        if self.headers.get("Notion-Version") != NOTION_VERSION:
            return self._json(400, {"object": "error", "code": "missing_version", "message": "Notion-Version header failed validation"})
        rest = url.path[len("/notion/v1/"):]
        q = {k: v[0] for k, v in parse_qs(url.query).items()}
        cursor = int((body.get("start_cursor") or q.get("start_cursor") or "0"))

        def page_of(items):
            more = cursor + SLACK_PAGE < len(items)
            return self._json(200, {"object": "list", "results": items[cursor:cursor + SLACK_PAGE], "has_more": more,
                                    "next_cursor": str(cursor + SLACK_PAGE) if more else None})

        def page_json(pid, spec):
            return {"object": "page", "id": pid, "url": f"https://www.notion.so/{pid.replace('-', '')}", "in_trash": spec.get("in_trash", False),
                    "last_edited_time": spec.get("edited", "2026-10-01T10:00:00.000Z"), "parent": spec.get("parent", {"type": "workspace", "workspace": True}),
                    "properties": {"title": {"type": "title", "title": [{"plain_text": spec["title"]}]}}}
        if method == "GET" and rest == "users/me":
            return self._json(200, {"object": "user", "type": "bot", "bot": {"workspace_name": NOTION["workspace"]}})
        if method == "POST" and rest == "search":
            kind = ((body.get("filter") or {}).get("value"))
            if kind == "data_source":
                return page_of([{"object": "data_source", "id": sid, "title": [{"plain_text": s["title"]}]}
                                for sid, s in NOTION["sources"].items() if s.get("shared")])
            rows = {r for s in NOTION["sources"].values() for r in s.get("rows", [])}
            return page_of([page_json(pid, spec) for pid, spec in NOTION["pages"].items() if spec.get("shared") and pid not in rows])
        parts = rest.split("/")
        if method == "POST" and len(parts) == 3 and parts[0] == "data_sources" and parts[2] == "query":
            source = next((s for sid, s in NOTION["sources"].items() if sid.replace("-", "") == parts[1].replace("-", "")), None)
            if not source or not source.get("shared"):
                return self._json(404, {"object": "error", "code": "object_not_found"})
            return page_of([page_json(pid, NOTION["pages"][pid]) for pid in source.get("rows", [])])
        if method == "GET" and len(parts) == 3 and parts[0] == "blocks" and parts[2] == "children":
            want = parts[1].replace("-", "")  # Notion takes an id with or without its dashes
            spec = next((s for pid, s in NOTION["pages"].items() if pid.replace("-", "") == want), None)                 or next((b for p in NOTION["pages"].values() for b in p.get("blocks", []) if str(b.get("id")).replace("-", "") == want), None)
            if spec is None:
                return self._json(404, {"object": "error", "code": "object_not_found"})
            return page_of(spec.get("blocks") or spec.get("children") or [])
        self._json(404, {"object": "error", "code": "invalid_request_url"})

    def _slack(self, url, state: dict = SLACK):
        """The Web API methods the connector calls, at Slack's own paths, each answering as Slack does."""
        if self.headers.get("Authorization") != f"Bearer {SLACK_TOKEN}":
            return self._json(200, {"ok": False, "error": "invalid_auth"})
        method, q = url.path.split("/api/", 1)[1], {k: v[0] for k, v in parse_qs(url.query).items()}
        team, channels = state["team"], state["channels"]
        start = int((q.get("cursor") or "c:0")[2:])

        def page(key, items, **more):
            body = {"ok": True, key: items[start:start + SLACK_PAGE], **more}
            if start + SLACK_PAGE < len(items):
                body["response_metadata"] = {"next_cursor": f"c:{start + SLACK_PAGE}"}
            return self._json(200, body)
        channel = channels.get(q.get("channel", ""))
        if method == "auth.test":
            return self._json(200, {"ok": True, "team": team["name"], "team_id": team["id"], "user": "aerugo", "user_id": SLACK_BOT,
                                    "url": team["url"] + "/"})
        if method == "conversations.list":
            return page("channels", [{"id": cid, "name": c["name"], "is_private": c.get("private", False), "is_member": c.get("bot_member", True),
                                      "is_ext_shared": c.get("shared", False), "is_shared": c.get("shared", False)} for cid, c in channels.items()])
        if method == "users.list":
            return page("members", state["users"])
        if method == "users.info":  # one person (C2b: is the one asking an admin, an owner, this workspace's?)
            found = next((u for u in state["users"] if u["id"] == q.get("user")), None)
            return self._json(200, {"ok": True, "user": found} if found else {"ok": False, "error": "user_not_found"})
        if channel is None:
            return self._json(200, {"ok": False, "error": "channel_not_found"})
        if method == "conversations.info":
            return self._json(200, {"ok": True, "channel": {"id": q["channel"], "name": channel["name"], "is_private": channel.get("private", False),
                                                            "is_member": channel.get("bot_member", True), "creator": channel.get("creator"),
                                                            "is_ext_shared": channel.get("shared", False), "is_shared": channel.get("shared", False)}})
        if not channel.get("bot_member", True):
            return self._json(200, {"ok": False, "error": "not_in_channel"})
        if method == "conversations.members":
            return page("members", channel.get("members", []))
        if method == "reactions.get":  # a message's reactions, a reply's too (C2b's "don't capture")
            found = next((m for top in channel.get("messages", []) for m in [top, *top.get("replies", [])] if m["ts"] == q.get("timestamp")), None)
            return self._json(200, {"ok": True, "type": "message", "message": {k: v for k, v in found.items() if k != "replies"}}
                              if found else {"ok": False, "error": "message_not_found"})
        if method == "conversations.history":
            if state.get("throttle"):
                state["throttle"] = False
                self.send_response(429)
                self.send_header("Retry-After", "1")
                self.send_header("Content-Length", "0")
                self.end_headers()
                return
            oldest = float(q.get("oldest") or 0)
            if q.get("cursor") is None:
                state.setdefault("history_calls", []).append({"channel": q.get("channel"), "oldest": oldest})  # H6
            latest = float(q.get("latest") or "inf")
            top = sorted((m for m in channel.get("messages", []) if oldest < float(m["ts"]) < latest), key=lambda m: -float(m["ts"]))
            return page("messages", [{k: v for k, v in m.items() if k != "replies"} for m in top], has_more=False)
        if method == "conversations.replies":
            parent = next((m for m in channel.get("messages", []) if m["ts"] == q.get("ts")), None)
            if parent is None:
                return self._json(200, {"ok": False, "error": "thread_not_found"})
            return page("messages", [{k: v for k, v in parent.items() if k != "replies"}] + parent.get("replies", []), has_more=False)
        self._json(200, {"ok": False, "error": "unknown_method"})

    def _graph(self, url):
        """Graph's users, deleted users, groups, and transitive members, and SharePoint's sites, libraries, files, and
        permissions: small pages, each with an absolute @odata.nextLink (or, when smoke asks, one on another host,
        which a client must never follow)."""
        auth = self.headers.get("Authorization") or ""
        app = auth[len(f"Bearer {GRAPH_TOKEN}."):] if auth.startswith(f"Bearer {GRAPH_TOKEN}.") else None
        if app not in GRAPH["apps"]:
            return self._json(401, {"error": {"code": "InvalidAuthenticationToken", "message": "Access token is empty."}})
        q = {k: v[0] for k, v in parse_qs(url.query).items()}
        rest = url.path[len("/microsoft/v1.0/"):]
        start = int(q.get("$skiptoken") or 0)

        def page(items, delta_link=None):
            body = {"value": items[start:start + ADMIN_PAGE]}
            if start + ADMIN_PAGE < len(items):
                host = "http://elsewhere.test/v1.0" if GRAPH["next_elsewhere"] else f"{MICROSOFT}/v1.0"
                from urllib.parse import urlencode  # Graph's next link keeps the request's own query, $count included
                body["@odata.nextLink"] = f"{host}/{rest}?{urlencode({**q, '$skiptoken': start + ADMIN_PAGE}, safe='$,')}"
            elif delta_link:  # the last page of a delta says where the next one starts (H6)
                body["@odata.deltaLink"] = delta_link
            return self._json(200, body)
        parts = rest.split("/")
        if parts[0] == "sites" and len(parts) in (2, 3):  # a site, and its libraries: only to an app it's granted to
            site = GRAPH["sites"].get(parts[1])
            if site is None or app not in site.get("granted", []):
                return self._json(403 if site else 404, {"error": {"code": "accessDenied", "message": "Access denied"}})
            if len(parts) == 2:
                return self._json(200, {"id": parts[1], "displayName": site["name"]})
            if parts[2] == "drives":
                return page([{"id": did, "name": d["name"], "driveType": "documentLibrary"} for did, d in site["drives"].items()])
        if parts[0] == "drives" and len(parts) >= 3:
            found = next(((s, s["drives"][parts[1]]) for s in GRAPH["sites"].values() if parts[1] in s["drives"]), None)
            if found is None or app not in found[0].get("granted", []):
                return self._json(403, {"error": {"code": "accessDenied", "message": "Access denied"}})
            drive = found[1]
            if parts[2:] == ["root", "delta"]:
                state = DELTA.get(parts[1], {"seq": 0, "changed": {}, "gone": {}})
                token = q.get("token")
                if start == 0:
                    DELTA_CALLS.append({"drive": parts[1], "token": token})
                if token is not None and int(token) <= state.get("expired_to", -1):  # as Graph answers an expired link (H6b)
                    return self._json(410, {"error": {"code": "resyncRequired", "message": "Resync required. Replace any "
                                                      "local items with the server's version and start over."}})
                if token is None:  # from the start: everything now
                    listed = [{"id": "root", "name": "root", "root": {}, "folder": {}}] + drive["items"]
                else:
                    listed = ([i for i in drive["items"] if state["changed"].get(i["id"], 0) > int(token)]
                              + [{"id": iid, "deleted": {}} for iid, at in state["gone"].items() if at > int(token)])
                return page(listed, f"{MICROSOFT}/v1.0/drives/{parts[1]}/root/delta?token={state['seq']}")
            item = next((i for i in drive["items"] if len(parts) >= 4 and i["id"] == parts[3]), None)
            if item is None:
                return self._json(404, {"error": {"code": "itemNotFound"}})
            if len(parts) == 4:  # an item itself, as a delta's changed file asks for the folders above it
                return self._json(200, {k: v for k, v in item.items() if k not in ("text", "permissions")})
            if parts[4:] == ["permissions"]:
                return page(item.get("permissions", []))
            if parts[4:] == ["content"]:
                where = "http://elsewhere.test/download" if item.get("bad_download") else f"{MICROSOFT}/download/{parts[1]}/{parts[3]}"
                self.send_response(302)
                self.send_header("Location", where)
                self.send_header("Content-Length", "0")
                self.end_headers()
                return
        if parts[0] == "teams" and len(parts) >= 2:  # Teams (C1a): only for an app the team's owner consented to
            team = GRAPH["teams"].get(parts[1])
            if team is None or app not in team.get("consented", []):
                return self._json(403 if team else 404, {"error": {"code": "Forbidden", "message": "Missing role permissions"}})
            if len(parts) == 2:
                return self._json(200, {"id": parts[1], "displayName": team["name"]})
            if parts[2:] == ["members"]:
                return page(team.get("members", []))
            if parts[2:] == ["channels"]:
                return page([{"id": cid, "displayName": ch["name"], "membershipType": ch.get("membershipType", "standard")}
                             for cid, ch in team["channels"].items()])
            channel = team["channels"].get(parts[3]) if len(parts) >= 5 and parts[2] == "channels" and parts[4] == "messages" else None
            if channel is None:
                return self._json(404, {"error": {"code": "NotFound"}})
            posts = channel.get("messages", [])
            bare = lambda m: {k: v for k, v in m.items() if k != "replies"}
            if len(parts) == 5:
                return page([bare(m) for m in posts])
            post = next((m for m in posts if m["id"] == parts[5]), None)
            if post is None:
                return self._json(404, {"error": {"code": "NotFound"}})
            if len(parts) == 6:
                return self._json(200, bare(post))
            if parts[6:] == ["replies"]:
                return page(post.get("replies", []))
            reply = next((r for r in post.get("replies", []) if len(parts) == 8 and parts[6] == "replies" and r["id"] == parts[7]), None)
            return self._json(200, reply) if reply else self._json(404, {"error": {"code": "NotFound"}})
        if rest == "users":
            return page(GRAPH["users"])
        if rest == "directory/deletedItems/microsoft.graph.user":
            return page(GRAPH["deleted"])
        if rest == "groups":
            return page([{"id": gid, "mail": g.get("mail"), "displayName": g.get("name")} for gid, g in GRAPH["groups"].items()])
        parts = rest.split("/")
        if len(parts) == 4 and parts[0] == "groups" and parts[2] == "transitiveMembers" and parts[3] == "microsoft.graph.user":
            if self.headers.get("ConsistencyLevel") != "eventual" or q.get("$count") != "true":  # as Graph insists for a cast
                return self._json(400, {"error": {"code": "Request_UnsupportedQuery", "message": "requires ConsistencyLevel and $count"}})
            if parts[1] not in GRAPH["groups"]:
                return self._json(404, {"error": {"code": "Request_ResourceNotFound"}})
            return page([{k: u.get(k) for k in ("id", "mail", "userPrincipalName")} for u in graph_transitive(parts[1])])
        self._json(404, {"error": {"code": "BadRequest", "message": "not found"}})

    def _admin(self, url):
        """The Directory API's users, groups, and members, as Google pages them (pageToken, never a URL)."""
        if self.headers.get("Authorization") != f"Bearer {ADMIN_TOKEN}":
            return self._json(401, {"error": {"code": 401, "message": "Login Required", "status": "UNAUTHENTICATED"}})
        if ADMIN.get("broken"):  # a listing that can't finish: the client must fail, not read it as everyone leaving
            return self._json(503, {"error": {"code": 503, "message": "Backend Error"}})
        q = {k: v[0] for k, v in parse_qs(url.query).items()}
        rest = url.path[len("/google/admin/directory/v1/"):]
        page = lambda items, key: self._json(200, {key: items[int(q.get("pageToken") or 0):int(q.get("pageToken") or 0) + ADMIN_PAGE],
                                                   **({"nextPageToken": str(int(q.get("pageToken") or 0) + ADMIN_PAGE)}
                                                      if int(q.get("pageToken") or 0) + ADMIN_PAGE < len(items) else {})})
        if rest == "users" and q.get("customer") == "my_customer":
            return page(ADMIN["deleted"] if q.get("showDeleted") == "true" else ADMIN["users"], "users")
        if rest == "groups" and q.get("customer") == "my_customer":
            return page([{"email": g} for g in ADMIN["groups"]], "groups")
        parts = rest.split("/")
        if len(parts) == 3 and parts[0] == "groups" and parts[2] == "members":
            from urllib.parse import unquote
            group = unquote(parts[1])
            if group not in ADMIN["groups"]:
                return self._json(404, {"error": {"code": 404, "message": "Resource Not Found: groupKey"}})
            return page(admin_members(group, q.get("includeDerivedMembership") == "true"), "members")
        self._json(404, {"error": {"code": 404, "message": "not found"}})

    def _jira(self, url, site: dict):
        import base64
        if self.headers.get("Authorization") != "Basic " + base64.b64encode(JIRA_AUTH.encode()).decode():
            return self._json(401, {"errorMessages": ["Client must be authenticated to access this resource."]})
        q, rest = {k: v[0] for k, v in parse_qs(url.query).items()}, url.path.split("/rest/api/3/", 1)[1]
        projects = site["projects"]
        doc = lambda text: {"type": "doc", "version": 1, "content": [{"type": "paragraph", "content": [{"type": "text", "text": text}]}]}
        if rest == "myself":
            return self._json(200, {"accountId": "aerugo-bot", "displayName": "Aerugo bot"})
        if rest == "project/search":
            return self._json(200, {"values": [{"id": k, "key": k, "name": p["name"]} for k, p in projects.items()], "isLast": True})
        as_issue = lambda i: {"id": i["key"], "key": i["key"], "fields": {
            "summary": i["summary"], "description": doc(i.get("description", "")), "updated": i.get("updated", "2026-10-01T10:00:00.000+0000"),
            "security": i.get("security"), "issuetype": {"name": "Task"}, "status": {"name": "To Do"}, "labels": [],
            "assignee": i.get("assignee"), "reporter": {"displayName": "Smoke"},
            "resolution": i.get("resolution"), "resolutiondate": i.get("resolutiondate"),
            "attachment": [{"id": a["id"], "filename": a["filename"], "size": len(a["text"]),
                            "content": f"http://{self.headers.get('Host')}/{url.path.split('/')[1]}/rest/api/3/attachment/content/{a['id']}"} for a in i.get("attachments", [])],
            "comment": {"comments": [{"author": {"displayName": c["author"]}, "created": c.get("created", "2026-10-01T10:00:00.000+0000"),
                                      "body": doc(c["body"]),  # a group or role restriction, and Service Management's internal notes, as Jira sends them
                                      **{k: c[k] for k in ("visibility", "jsdPublic") if k in c}} for c in i.get("comments", [])]}}}
        if rest == "search/jql":  # project = KEY, and when a sync asks for changes only (H6), AND updated >= -Nm
            from datetime import datetime, timedelta, timezone
            jql = q.get("jql", "")
            site.setdefault("searches", []).append(jql)
            key = (re.search(r"project = (\S+)", jql) or [None, ""])[1]
            terms = re.findall(r'\b(summary|description|text) ~ "([a-z0-9]+)"', jql)
            if terms:  # Jira's own search (F11b), as a live-only connection asks it: any term, most words first
                listed = re.search(r"project in \(([^)]*)\)", jql)
                keys = [k.strip() for k in listed[1].split(",")] if listed else list(projects)
                # each field as Jira matches it; `text` is every text field, every comment the account sees included
                fields = lambda i: {"summary": i["summary"], "description": i.get("description", ""), "text": " ".join(
                    [i["summary"], i.get("description", "")] + [c["body"] for c in i.get("comments", [])])}
                words = lambda i: {f: set(re.findall(r"[a-z0-9]+", t.lower())) for f, t in fields(i).items()}
                hits = [(len({w for f, w in terms if w in words(i)[f]}), n, i) for n, i in
                        enumerate(i for k in keys for i in projects.get(k, {}).get("issues", []))]
                found = [i for score, _, i in sorted(hits, key=lambda h: (-h[0], h[1])) if score]
                return self._json(200, {"issues": [{"id": i["key"], "key": i["key"], "fields": {"summary": i["summary"]}}
                                                   for i in found[:int(q.get("maxResults") or 50)]], "isLast": True})
            within = re.search(r"updated >= -(\d+)m", jql)
            recent = lambda i: not within or datetime.strptime(i.get("updated", "2026-10-01T10:00:00.000+0000"), "%Y-%m-%dT%H:%M:%S.%f%z") \
                >= datetime.now(timezone.utc) - timedelta(minutes=int(within[1]))
            issues = [as_issue(i) for i in projects.get(key, {}).get("issues", []) if recent(i)]
            return self._json(200, {"issues": issues, "isLast": True})
        parts = rest.split("/")
        if len(parts) == 2 and parts[0] == "issue":  # one issue: a pointer's text, read live (F11a), each read recorded
            found = next((i for p in projects.values() for i in p.get("issues", []) if i["key"] == parts[1]), None)
            site.setdefault("issue_reads", []).append(parts[1])
            return self._json(200, as_issue(found)) if found else self._json(404, {"errorMessages": ["Issue does not exist"]})
        if len(parts) == 3 and parts[:2] == ["attachment", "content"]:  # an attachment's bytes (Q7)
            time.sleep(min(float(site.get("delay") or 0), 5))  # a slow site, when smoke asks: a sync then takes a while (C1)
            found = next((a for p in projects.values() for i in p.get("issues", []) for a in i.get("attachments", []) if a["id"] == parts[2]), None)
            return self._bytes(found["text"].encode()) if found else self._json(404, {})
        if len(parts) == 3 and parts[0] == "project" and parts[2] == "permissionscheme" and parts[1] in projects:
            holder = "applicationRole" if projects[parts[1]].get("company_wide") else "projectRole"
            return self._json(200, {"permissions": [{"permission": "BROWSE_PROJECTS", "holder": {"type": holder}}]})
        if rest == "user/permission/search":
            if q.get("issueKey"):
                issue = next((i for p in projects.values() for i in p.get("issues", []) if i["key"] == q["issueKey"]), None)
                if issue is None:
                    return self._json(404, {"errorMessages": ["Issue does not exist"]})
                key = issue["key"].rsplit("-", 1)[0]
                users = issue.get("viewers") if issue.get("security") else projects[key].get("browsers", [])
            else:
                users = projects.get(q.get("projectKey", ""), {}).get("browsers", [])
            return self._json(200, [{"active": True, **u} for u in users or []])
        self._json(404, {"errorMessages": ["not found"]})

    def do_DELETE(self):
        if self.path == "/requests":  # a clean slate: both logs, and honest answers again
            REQUESTS.clear()
            RERANKS.clear()
            MODE["answer"], MODE["delay"] = ANSWER, 0.0
            return self._json(200, {})
        self._json(404, {})

    def do_POST(self):
        path = self.path.rstrip("/")
        if path == "/throttle":
            body = self._body()
            with THROTTLE_LOCK:
                THROTTLE.update({"prefix": str(body.get("prefix") or ""), "left": int(body.get("count") or 0), "answered": 0,
                                 "status": int(body.get("status") or 429), "retry_after": body.get("retry_after")})
            return self._json(200, {})
        if self._throttled():
            return None
        if path == "/alerts":  # the test webhook: the app posts alerts here when an org admin points it here
            ALERTS.append(self._body())
            return self._json(200, {})
        if path == "/oidc/token":  # a CI job asking for its workload token
            return self._json(200, {"token": ci_token(self._body())})
        if path == "/google/token":  # OAuth, as Google's token endpoint does
            form = parse_qs(self.rfile.read(int(self.headers.get("Content-Length", 0))).decode())
            if form.get("grant_type", [""])[0] == "urn:ietf:params:oauth:grant-type:jwt-bearer":  # a service account (ADR 0011)
                return self._json(*service_account_token(form.get("assertion", [""])[0]))
            if form.get("grant_type", [""])[0] == "authorization_code":  # sign-in: a one-time code for an ID token
                who = SIGNIN_CODES.pop(form.get("code", [""])[0], None)
                if not who:
                    return self._json(400, {"error": "invalid_grant"})
                return self._json(200, {"id_token": google_id_token(form.get("client_id", [""])[0], who),
                                        "access_token": "fake-signin-token", "token_type": "Bearer"})
            if form.get("refresh_token", [""])[0] != DRIVE_REFRESH:  # Drive: a refresh token for an access token
                return self._json(400, {"error": "invalid_grant"})
            return self._json(200, {"access_token": DRIVE_TOKEN, "expires_in": 3600, "token_type": "Bearer"})
        if path == "/google/signin":  # smoke signs someone in: the code Google would send back to the callback
            code = secrets.token_urlsafe(24)
            SIGNIN_CODES[code] = self._body()
            return self._json(200, {"code": code})
        if path == "/google/service-account":  # smoke mints a service account's JSON key
            return self._json(200, new_service_account(self._body()["email"]))
        if path == "/google/folders":  # smoke sets folders, and which are shared with the service account
            FOLDERS.clear()
            FOLDERS.update(self._body()["folders"])
            return self._json(200, {})
        if path == "/secret-stores":  # smoke sets the fake Vault's, AWS's, and Azure's secrets
            for store, values in self._body().items():
                SECRET_STORES[store].update(values)
            return self._json(200, {})
        if path == "/hcvault/v1/auth/jwt/login":  # Vault's JWT auth: the runner's identity token for the vault audience
            body = self._body()
            if body.get("role") != "aerugo" or body.get("jwt") != "fake-gcp-identity-token-for-vault":
                return self._json(403, {"errors": ["permission denied"]})
            return self._json(200, {"auth": {"client_token": VAULT_TOKEN, "lease_duration": 3600}})
        if path.startswith("/azure-login/") and path.endswith("/oauth2/v2.0/token"):  # Entra, by workload identity federation
            form = {k: v[0] for k, v in parse_qs(self.rfile.read(int(self.headers.get("Content-Length", 0))).decode()).items()}
            if form.get("client_assertion") != "fake-gcp-identity-token-for-api://AzureADTokenExchange" \
                    or form.get("scope") != "https://vault.azure.net/.default":
                return self._json(401, {"error": "invalid_client"})
            return self._json(200, {"access_token": AZURE_KV_TOKEN, "token_type": "Bearer", "expires_in": 3600})
        if path == "/catalog/setup":  # smoke sets the fake catalogs' tables, and Entra's object ids for Purview
            body = self._body()
            CATALOG.update({"tables": body.get("tables", []), "people": body.get("people", {})})
            return self._json(200, {"project": CATALOG_PROJECT, "databricks": DATABRICKS, "role": AWS["role"], "datahub_token": DATAHUB_TOKEN})
        if path == "/datahub/setup":  # kl.enterprise_eval sets a DataHub platform's tables, in CATALOG's shape
            DATAHUB_PLATFORMS.update(self._body().get("platforms", {}))
            return self._json(200, {"datahub_token": DATAHUB_TOKEN})
        if path.startswith("/purview/") and path.endswith("/oauth2/v2.0/token"):  # Entra's token endpoint, for Purview's audience
            form = parse_qs(self.rfile.read(int(self.headers.get("Content-Length", 0))).decode())
            return self._json(*entra_token(path.split("/")[2].lower(), form, PURVIEW, "https://purview.azure.net/.default", PURVIEW_TOKEN))
        if path.startswith(("/gcp-catalog/", "/databricks/", "/aws/", "/purview/", "/datahub/")):
            return self._catalog("POST", urlparse(self.path), self.rfile.read(int(self.headers.get("Content-Length", 0))))
        if path == "/gcp/secrets":  # smoke sets the fake Secret Manager's secrets, each a JSON object
            GCP_SECRETS.update(self._body()["secrets"])
            return self._json(200, {})
        if path == "/notion/setup":  # smoke sets the fake Notion's pages and data sources, and which are shared
            body = self._body()
            NOTION.update({"pages": body.get("pages", {}), "sources": body.get("sources", {})})
            return self._json(200, {})
        if path.startswith("/notion/v1/"):
            return self._notion("POST", urlparse(self.path), self._body())
        if path.endswith("/setup") and path.split("/")[1] in SLACK_SITES:  # a fake Slack workspace's people and channels
            body, state = self._body(), SLACK_SITES[path.split("/")[1]]
            state.update({"users": body.get("users", []), "channels": body.get("channels", {}), "throttle": bool(body.get("throttle")),
                          "history_calls": state.get("history_calls", []), "posted": state.get("posted", [])})
            return self._json(200, {})
        if path.split("/")[-1] in ("chat.postMessage", "chat.postEphemeral") and path.split("/")[1] in SLACK_SITES:
            state = SLACK_SITES[path.split("/")[1]]  # the bot's writes: C2b's notice, and C2c's private reply
            if self.headers.get("Authorization") != f"Bearer {SLACK_TOKEN}":
                return self._json(200, {"ok": False, "error": "invalid_auth"})
            form = {k: v[0] for k, v in parse_qs(self.rfile.read(int(self.headers.get("Content-Length", 0))).decode()).items()}
            if form.get("channel") not in state["channels"]:
                return self._json(200, {"ok": False, "error": "channel_not_found"})
            private = path.endswith("chat.postEphemeral")
            if private and form.get("user") not in state["channels"][form["channel"]].get("members", []):
                return self._json(200, {"ok": False, "error": "user_not_in_channel"})
            ts = f"{time.time():.6f}"
            state.setdefault("posted", []).append({"channel": form["channel"], "text": form.get("text", ""), "ts": ts,
                                                   **({"user": form["user"], "ephemeral": True, "thread": form.get("thread_ts")}
                                                      if private else {})})
            return self._json(200, {"ok": True, "channel": form["channel"], **({"message_ts": ts} if private else {"ts": ts})})
        if path == "/microsoft/app":  # smoke mints an app registration and its certificate
            body = self._body()
            return self._json(200, new_entra_app(body["tenant_id"], body.get("client_id")))
        if path == "/microsoft/teams":  # smoke sets Teams' teams, who consented, members, channels, and posts (C1a)
            GRAPH["teams"] = self._body()["teams"]
            return self._json(200, {})
        if path == "/botframework/sign":  # smoke makes a token as Bot Framework signs one for a bot (C1a)
            return self._json(200, {"token": bot_token(self._body())})
        if path == "/teams-bot/v3/conversations":  # the bot starts a post in a channel: capture's notice (C1a)
            auth = self.headers.get("Authorization") or ""
            app = auth[len(f"Bearer {BOT_TOKEN}."):] if auth.startswith(f"Bearer {BOT_TOKEN}.") else None
            body = self._body()
            if app not in GRAPH["apps"]:
                return self._json(401, {"error": {"code": "Unauthorized", "message": "Authorization has been denied"}})
            channel = ((body.get("channelData") or {}).get("channel") or {}).get("id")
            if not body.get("isGroup") or not any(channel in t["channels"] and app in t.get("consented", []) for t in GRAPH["teams"].values()):
                return self._json(403, {"error": {"code": "BotNotInConversationRoster"}})
            TEAMS_BOT["posted"].append({"channel": channel, "text": (body.get("activity") or {}).get("text", ""), "app": app})
            return self._json(201, {"id": f"{channel};messageid={int(time.time() * 1000)}", "activityId": str(len(TEAMS_BOT["posted"]))})
        if path == "/microsoft/sites":  # smoke sets SharePoint's sites, which apps each is granted to, and their files
            GRAPH["sites"] = self._body()["sites"]
            for site in GRAPH["sites"].values():
                for did, d in site["drives"].items():
                    state = DELTA.setdefault(did, {"seq": 0, "changed": {}, "gone": {}, "last": {}})
                    now = {i["id"]: i for i in d["items"]}
                    for iid, item in now.items():
                        if state["last"].get(iid) != item:
                            state["seq"] += 1
                            state["changed"][iid] = state["seq"]
                    for iid in set(state["last"]) - set(now):
                        state["seq"] += 1
                        state["gone"][iid] = state["seq"]
                    state["last"] = now
            return self._json(200, {})
        if path == "/microsoft/expire":  # H6b: every delta link a library has handed out so far is expired
            state = DELTA.setdefault(self._body()["drive"], {"seq": 0, "changed": {}, "gone": {}, "last": {}})
            state["expired_to"] = state["seq"]
            return self._json(200, {})
        if path == "/google/expire":  # H6b: every changes token a shared drive has handed out so far is refused
            state = DRIVE_CHANGES.setdefault(self._body()["drive"], {"seq": 0, "changed": {}, "gone": {}, "last": {}})
            state["expired_to"] = state["seq"]
            return self._json(200, {})
        if path == "/microsoft/setup":  # smoke sets the tenant's users, deleted users, and groups
            body = self._body()
            GRAPH.update({"users": body.get("users", []), "deleted": body.get("deleted", []), "groups": body.get("groups", {}),
                          "next_elsewhere": bool(body.get("next_elsewhere"))})
            GRAPH_REQUESTS.clear()
            return self._json(200, {})
        if path.startswith("/microsoft/") and path.endswith("/oauth2/v2.0/token"):  # Entra's token endpoint
            form = parse_qs(self.rfile.read(int(self.headers.get("Content-Length", 0))).decode())
            if form.get("scope", [""])[0] == BOT_SCOPE:  # Teams' bot, as its app (C1a)
                return self._json(*entra_token(path.split("/")[2].lower(), form, scope=BOT_SCOPE, token=BOT_TOKEN))
            return self._json(*entra_token(path.split("/")[2].lower(), form))
        if path == "/google/admin/setup":  # smoke sets the fake Workspace directory, and which admins delegation allows
            ADMIN.clear()
            ADMIN.update({"users": [], "deleted": [], "groups": {}, "subjects": [], "broken": False, **self._body()})
            TOKEN_REQUESTS.clear()
            return self._json(200, {})
        if path == "/okta/setup":  # smoke sets the fake Okta's users
            OKTA["users"] = self._body()["users"]
            return self._json(200, {})
        if path.endswith("/setup") and path.split("/")[1] in CONFLUENCE_SITES:  # a fake Confluence site's spaces, pages, groups, and people
            state = CONFLUENCE_SITES[path.split("/")[1]]
            listings = state.get("listings", [])
            state.clear()
            state.update({"spaces": {}, "groups": {}, "users": {}, **self._body(), "listings": listings})
            return self._json(200, {})
        if path.endswith("/setup") and path.split("/")[1] in JIRA_SITES:  # a fake Jira site's projects, issues, and who may see each
            site = JIRA_SITES[path.split("/")[1]]
            site.clear()
            site.update({"projects": {}, **self._body()})
            return self._json(200, {})
        if path == "/google/files":  # smoke sets a shared drive's files (and who each is shared with)
            request = self._body()
            DRIVES[request["drive"]] = {"name": request.get("name", request["drive"]), "files": {f["id"]: f for f in request["files"]}}
            state = DRIVE_CHANGES.setdefault(request["drive"], {"seq": 0, "changed": {}, "gone": {}, "last": {}})
            now = DRIVES[request["drive"]]["files"]
            for fid, f in now.items():
                if state["last"].get(fid) != f:
                    state["seq"] += 1
                    state["changed"][fid] = state["seq"]
            for fid in set(state["last"]) - set(now):
                state["seq"] += 1
                state["gone"][fid] = state["seq"]
            state["last"] = dict(now)
            return self._json(200, {})
        if path == "/mode":
            body = self._body()
            MODE["answer"] = UNFAITHFUL if body.get("answer") == "unfaithful" else ANSWER
            MODE["delay"] = min(float(body.get("delay") or 0), 10)
            return self._json(200, {})
        if path == "/v1/rerank":
            time.sleep(MODE["delay"])
            request = self._body()
            RERANKS.append(request)
            asked = words(request.get("query", ""))
            docs = request.get("documents") or []
            results = [{"index": i, "relevance_score": len(asked & words(d)) / (1 + len(asked))} for i, d in enumerate(docs)]
            results.sort(key=lambda r: r["relevance_score"], reverse=True)
            return self._json(200, {"results": results[: request.get("top_n") or len(results)]})
        if path == "/v1/audio/transcriptions":  # a speech model (Q8): a fixed transcript, with timestamps
            REQUESTS.append({"transcribe": self.headers.get("Content-Type", "")[:40]})
            self.rfile.read(int(self.headers.get("Content-Length", 0)))
            return self._json(200, {"text": "Standup.", "segments": [{"start": 0.0, "text": "Standup: the cutover rehearsal went well."},
                                                                    {"start": 75.0, "text": "Next, we fail over at noon on Friday."}]})
        if path != "/v1/chat/completions":
            return self._json(404, {})
        time.sleep(MODE["delay"])
        request = self._body()
        last = (request.get("messages") or [{}])[-1].get("content")
        if isinstance(last, list) and any(p.get("type") == "image_url" for p in last):  # a vision model (Q8)
            REQUESTS.append({"vision": True})
            return self._json(200, {"choices": [{"message": {"role": "assistant", "content": "A flowchart: Orders feed the Ledger, which feeds Settlement."}}],
                                    "usage": {"prompt_tokens": 80, "completion_tokens": 14}})
        REQUESTS.append(request)
        messages = request.get("messages") or []
        system = next((m.get("content", "") for m in messages if m.get("role") == "system"), "")
        prompt = messages[-1].get("content", "") if messages else ""
        if "knowledge compiler" in system:
            content = compiler_draft(prompt)
        elif "knowledge extractor" in system:
            content = extraction_reply(prompt)
        elif "grade answers" in system:
            content = grade(prompt)
        else:
            content = MODE["answer"]
        self._json(200, {"choices": [{"message": {"role": "assistant", "content": content}}],
                         "usage": {"prompt_tokens": 100, "completion_tokens": 12}})


def extraction_reply(prompt: str) -> str:
    """An extractor's reply (F2), mechanical on purpose: a fact per section, its first sentence cited, titled by its
    heading and the document; a section saying make-an-unsupported-unit also gets a sentence it doesn't state, so smoke
    can prove such a proposal is dropped. It measures the pipeline, never the quality of extraction."""
    title = re.search(r"^Document: (.+)$", prompt, re.M)
    units = []
    for n, heading, body in re.findall(r"^\[(\d+)\] ([^\n]*)\n(.*?)(?=\n\n\[\d+\] |\Z)", prompt, re.M | re.S):
        first = re.split(r"(?<=[.!?])\s", body.strip(), maxsplit=1)[0].strip()
        if len(first.split()) >= 3:
            units.append({"kind": "fact", "topic": heading, "title": f"{heading} ({title.group(1) if title else ''})"[:200],
                          "text": f"{first.rstrip('.')} [{n}]."})
        if "make-an-unsupported-unit" in body:
            units.append({"kind": "fact", "topic": heading, "title": f"{heading}: the move", "text": f"The queue moves to the 2031 platform on Friday [{n}]."})
    return json.dumps({"units": units})


def compiler_draft(prompt: str) -> str:
    """A compiler reply. A topic containing make-a-bad-draft cites a source the compiler wasn't given,
    so smoke can prove such drafts are rejected."""
    cite = "[^s9]" if "make-a-bad-draft" in prompt else "[^s1]"
    return json.dumps({"type": "Playbook", "title": "Paging during a payments outage",
                       "body": f"When payments are down, page the payments on-call lead first {cite}."})


def grade(prompt: str) -> str:
    """The eval's judge, in miniature: a claim is supported when most of its words appear in the cited
    sources, and a key fact is stated when most of its words appear in the answer."""
    sources, _, rest = prompt.partition("Answer to grade:")
    found = re.search(r"<<<\n(.*?)\n>>>", rest, re.S)
    answer = found.group(1) if found else ""
    facts = re.findall(r"^- (.+)$", rest.partition("Key facts to look for:")[2], re.M)
    pool, said = words(sources), words(answer)

    def covered(text: str, known: set[str]) -> bool:
        w = words(text)
        return bool(w) and len(w & known) >= 0.6 * len(w)

    claims = [c.strip() for c in re.split(r"(?<=[.!?])\s+", re.sub(r"\s*\[\d+\]", "", answer)) if c.strip()]
    return json.dumps({"claims": [{"claim": c, "supported": covered(c, pool)} for c in claims],
                       "facts": [covered(f, said) for f in facts]})


if __name__ == "__main__":
    ThreadingHTTPServer(("0.0.0.0", 8080), Handler).serve_forever()
